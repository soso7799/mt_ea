"""
事先定好的單一假設：在昨日與各時段高低點掛限價單做反轉（不做參數搜尋，完整報告所有結果）

  關卡：昨高、昨低、昨亞高/低、昨歐高/低、昨美高/低（8 條，每天用前一日資料）
  掛單：關卡在價格上方 → 賣出限價；在下方 → 買入限價；價格須從一側靠近（前一根收盤距關卡 >0.1 ATR）
        每條關卡每天只成交第一次
  成交：賣出限價在 Bid 最高價 ≥ 關卡時成交；買入限價在 Ask 最低價（Bid 低 + 點差）≤ 關卡時成交
  停損：關卡外 0.3 ATR；停利：往回 1.0 ATR；最多持有 24 小時後以收盤平倉
        賣單以 Ask（= Bid + 點差）判斷停損/停利、買單以 Bid 判斷；成交那根若也碰到停損 → 算停損
  濾網：不限 / 靠近時量縮（前一根已收盤 K 線成交量 < 前 20 根平均 × 0.8，下單前就已知）
  報酬單位：R（1R = 0.3 ATR 停損距離），已含點差並扣外匯手續費

用法：python limit_reversal.py --cache H:\\...\\export\\bars_mtf --out H:\\...\\reports [--tf M15|H1] [--offline]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from analyze import closed_bars                        # noqa: E402
from build_report import DEFAULT_SYMBOLS               # noqa: E402
from data import TF_MIN, fetch_all                     # noqa: E402
from indicators import atr                             # noqa: E402
from levels import server_to_utc                       # noqa: E402
from levels_study import daily_levels, is_fx           # noqa: E402

LEVELS = ["昨高", "昨低", "昨亞高", "昨亞低", "昨歐高", "昨歐低", "昨美高", "昨美低"]
SL_ATR, TP_ATR, TOL, HOLD_H, VOL_TH = 0.3, 1.0, 0.1, 24, 0.8


def run_symbol(sym, df, point, contract, commission):
    o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
    sp = df["spread"].to_numpy(float) * point
    sp = np.where(sp > 0, sp, np.nanmedian(sp[sp > 0]) if (sp > 0).any() else point)
    tv = df["tick_volume"].to_numpy(float)
    volr = tv / pd.Series(tv).rolling(20, min_periods=20).mean().shift(1).to_numpy()
    n = len(c)
    lv, day = daily_levels(df, server_to_utc(df["time"]).hour.to_numpy())
    a14 = atr(h, l, c, 14)
    hold = max(1, int(HOLD_H * 60 / (df["time"].diff().dt.total_seconds().median() / 60)))
    day_start = np.r_[0, np.flatnonzero(day[1:] != day[:-1]) + 1]
    day_end = np.r_[day_start[1:], n]
    out = []
    for name in LEVELS:
        L_all = lv[name]
        for a0, b0 in zip(day_start, day_end):
            L = L_all[a0]
            if np.isnan(L):
                continue
            for i in range(max(a0, 1), b0):
                at = a14[i - 1]
                if np.isnan(at) or at <= 0:
                    continue
                tol = TOL * at
                if c[i - 1] < L - tol:
                    d = -1                        # 關卡在上方 → 賣出限價
                    filled = h[i] >= L
                elif c[i - 1] > L + tol:
                    d = 1                         # 關卡在下方 → 買入限價（Ask 碰到）
                    filled = l[i] + sp[i] <= L
                else:
                    continue
                if not filled:
                    continue
                sl, tp = L - d * SL_ATR * at, L + d * TP_ATR * at
                res, exit_px, j_exit = None, None, None
                # 成交那根：只檢查停損（保守）
                adj = sp[i] if d < 0 else 0.0
                if (d < 0 and h[i] + adj >= sl) or (d > 0 and l[i] <= sl):
                    res, exit_px, j_exit = "SL", sl, i
                else:
                    for j in range(i + 1, min(n, i + 1 + hold)):
                        adj = sp[j] if d < 0 else 0.0
                        hi, lo = h[j] + adj, l[j] + adj
                        s_hit = lo <= sl if d > 0 else hi >= sl
                        t_hit = hi >= tp if d > 0 else lo <= tp
                        if s_hit:
                            res, exit_px, j_exit = "SL", sl, j
                            break
                        if t_hit:
                            res, exit_px, j_exit = "TP", tp, j
                            break
                    if res is None:
                        j_exit = min(n - 1, i + hold)
                        res, exit_px = "TIME", c[j_exit] + (sp[j_exit] if d < 0 else 0.0)
                comm = 0.0
                if is_fx(sym) and commission:
                    comm = commission / contract * (1.0 if sym.endswith("USD") else L)
                pnl = (exit_px - L) * d - comm
                out.append(dict(symbol=sym, time=df["time"].iloc[i], level=name, side="賣(壓力)" if d < 0 else "買(支撐)",
                                vol=volr[i - 1], quiet=bool(volr[i - 1] < VOL_TH) if not np.isnan(volr[i - 1]) else False,
                                result=res, R=pnl / (SL_ATR * at), cost_R=(sp[i] + comm) / (SL_ATR * at),
                                bars=j_exit - i))
                break                               # 每條關卡每天只做第一次
    return out


def summarize(g):
    r = g["R"].to_numpy()
    n = len(r)
    if n < 2:
        return dict(筆數=n)
    t = r.mean() / r.std(ddof=1) * np.sqrt(n) if r.std(ddof=1) > 0 else np.nan
    return dict(筆數=n, 勝率=(g["result"] == "TP").mean() * 100, 停損率=(g["result"] == "SL").mean() * 100,
                平均R=r.mean(), t值=t, 總R=r.sum(), 平均成本R=g["cost_R"].mean())


def main():
    ap = argparse.ArgumentParser(description="關卡限價反轉（事先定好的單一假設）")
    ap.add_argument("--symbols", nargs="*", default=DEFAULT_SYMBOLS)
    ap.add_argument("--cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("--tf", default="M15", choices=["M5", "M15", "H1"])
    ap.add_argument("--commission", type=float, default=5.0)
    args = ap.parse_args()

    frames, metas = fetch_all(args.symbols, args.cache, args.terminal, offline=args.offline, max_age_min=24 * 60)
    ticks = [pd.Timestamp(m["tick_time"]) for m in metas.values() if m.get("tick_time")]
    server_now = max(ticks) if ticks else None
    rows = []
    for s in args.symbols:
        if (s, args.tf) not in frames:
            continue
        meta = metas.get(s, {})
        df = closed_bars(frames[(s, args.tf)], args.tf, server_now)
        rows += run_symbol(s, df, float(meta.get("point", 0.0001)), float(meta.get("contract") or 100000),
                           args.commission)
    T = pd.DataFrame(rows).sort_values("time").reset_index(drop=True)
    cut = T["time"].iloc[len(T) // 2]
    T["half"] = np.where(T["time"] < cut, "前半", "後半")

    tabs = {}
    variants = {"不限": T, "靠近時量縮": T[T["quiet"]]}
    main_rows = []
    for vname, g in variants.items():
        for part, gg in [("全部", g), ("前半", g[g["half"] == "前半"]), ("後半", g[g["half"] == "後半"])]:
            main_rows.append(dict(濾網=vname, 期間=part, **summarize(gg)))
    tabs["總結"] = pd.DataFrame(main_rows)
    for key, col in [("各商品", "symbol"), ("各關卡", "level"), ("買賣方向", "side")]:
        rr = []
        for vname, g in variants.items():
            for k, gg in g.groupby(col):
                s1 = summarize(gg)
                s2 = summarize(gg[gg["half"] == "前半"])
                s3 = summarize(gg[gg["half"] == "後半"])
                rr.append(dict(濾網=vname, **{col: k}, **s1, 前半平均R=s2.get("平均R"), 後半平均R=s3.get("平均R")))
        tabs[key] = pd.DataFrame(rr)

    ok = []
    for vname, g in variants.items():
        s1, s2, s3 = summarize(g), summarize(g[g["half"] == "前半"]), summarize(g[g["half"] == "後半"])
        per = g.groupby("symbol")["R"].mean()
        share = (per > 0).mean() if len(per) else 0
        ok.append((vname, s1, s2, s3, share))

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"限價反轉_{args.tf}_{updated:%Y%m%d_%H%M}.xlsx")
    notes = [f"產生時間 {updated}；週期 {args.tf}；期間 {T['time'].min()} ~ {T['time'].max()}；前後半分界 {cut}",
             "事先定好的單一假設，不做參數搜尋；以下所有結果完整列出。",
             "關卡：昨高、昨低、昨亞/歐/美盤高低（8 條）。關卡在上方掛賣出限價、在下方掛買入限價，價格須從一側靠近，每條關卡每天只做第一次。",
             f"停損 = 關卡外 {SL_ATR} ATR；停利 = 往回 {TP_ATR} ATR；最多持有 {HOLD_H} 小時。",
             "成交以 Bid/Ask 判斷（買單需 Ask 碰到關卡；賣單停損/停利以 Ask 判斷）；成交當根若也碰到停損 → 算停損（保守）。",
             f"濾網『靠近時量縮』= 前一根已收盤 K 線成交量 < 前 20 根平均 × {VOL_TH}（下單前已知）。",
             f"R = 損益 ÷ ({SL_ATR} ATR)，已含點差並扣外匯手續費 {args.commission} USD/手。損益兩平勝率約 {SL_ATR / (SL_ATR + TP_ATR) * 100:.0f}%（未計成本）。",
             "判定：前半、後半平均 R 都 >0、全部 t≥2、且 ≥60% 商品平均 R >0 → 值得寫成 EA。"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": notes}).to_excel(xw, sheet_name="說明", index=False)
        for k, v in tabs.items():
            v.round(3).to_excel(xw, sheet_name=k, index=False)
        T.round(4).to_excel(xw, sheet_name="交易明細", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.sheets.values():
            for cc in ws[1]:
                cc.font = Font(bold=True, color="FFFFFF")
                cc.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.5), 80)
            ws.freeze_panes = "A2"
    print(f"\n完成：{path}\n")
    for vname, s1, s2, s3, share in ok:
        verdict = (s2.get("平均R", -1) > 0 and s3.get("平均R", -1) > 0 and s1.get("t值", 0) >= 2 and share >= 0.6)
        print(f"[{vname}] 筆數 {s1.get('筆數')}  勝率 {s1.get('勝率', 0):.1f}%  平均 {s1.get('平均R', 0):+.3f}R  "
              f"t={s1.get('t值', 0):.2f}  前半 {s2.get('平均R', 0):+.3f}R  後半 {s3.get('平均R', 0):+.3f}R  "
              f"獲利商品 {share:.0%}  成本 {s1.get('平均成本R', 0):.2f}R  → {'✅ 值得寫成 EA' if verdict else '❌ 不成立'}")


if __name__ == "__main__":
    main()
