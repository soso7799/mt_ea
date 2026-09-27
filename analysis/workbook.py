"""把各分頁資料寫成 Excel；儀表板用下拉選單 + INDEX/MATCH 公式，換商品即更新。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ORDER = ["儀表板總表", "關卡", "多商品狀態總表", "多商品參數優化", "多商品時段壓力支撐", "多商品回測摘要",
         "多商品交易明細", "多商品進場信號", "波段避險分組", "總整理", "規則定義", "_dash"]
HDR = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE = Font(bold=True, size=14)
GOOD = PatternFill("solid", fgColor="C6EFCE")
BAD = PatternFill("solid", fgColor="FFC7CE")
WARN = PatternFill("solid", fgColor="FFEB9C")
TFS = ["D1", "H4", "H1", "M15", "M5"]


def _header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, maxw=45):
    w = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            v = str(c.value)
            ln = sum(2 if ord(ch) > 255 else 1 for ch in v)
            w[c.column_letter] = max(w.get(c.column_letter, 0), min(ln, maxw))
    for k, v in w.items():
        ws.column_dimensions[k].width = max(8, v + 2)


def _color(ws, col_name, rules):
    hdr = [c.value for c in ws[1]]
    if col_name not in hdr:
        return
    ci = hdr.index(col_name) + 1
    for r_ in range(2, ws.max_row + 1):
        c = ws.cell(row=r_, column=ci)
        for test, fill in rules:
            if c.value is not None and test(c.value):
                c.fill = fill
                break


def write_workbook(path, sheets, symbols, dash_symbol):
    wb = Workbook()
    wb.remove(wb.active)
    for name in ORDER:
        ws = wb.create_sheet(name)
        if name == "儀表板總表":
            continue
        for row in sheets[name]:
            ws.append(row)
        if name == "波段避險分組":
            ws["A1"].font = TITLE
            _header(ws, 5, 6)
        elif name != "_dash":
            _header(ws, 1, len(sheets[name][0]))
            ws.freeze_panes = "C2" if name != "規則定義" else "A2"
        for r_ in ws.iter_rows():
            for c in r_:
                if hasattr(c.value, "year") and hasattr(c.value, "hour"):
                    c.number_format = "yyyy-mm-dd hh:mm"
        _widths(ws, 80 if name == "規則定義" else 45)

    # 顏色標示
    _color(wb["多商品參數優化"], "判定", [(lambda v: v == "樣本外獲利", GOOD), (lambda v: v == "樣本外虧損", BAD)])
    _color(wb["多商品回測摘要"], "AvgReturnPct", [(lambda v: v > 0, GOOD), (lambda v: v <= 0, BAD)])
    _color(wb["多商品回測摘要"], "後30%AvgReturnPct", [(lambda v: v > 0, GOOD), (lambda v: v <= 0, BAD)])
    for n in ("多商品進場信號", "總整理"):
        _color(wb[n], "規則等級", [(lambda v: v == "可進場", GOOD), (lambda v: v == "觀察", WARN),
                                  (lambda v: v == "排除", BAD)])
    for n in ("多商品狀態總表",):
        _color(wb[n], "Status", [(lambda v: v == "多頭確認", GOOD), (lambda v: v == "空頭確認", BAD)])

    # 儀表板
    ws = wb["儀表板總表"]
    d = wb["_dash"]
    d.sheet_state = "hidden"
    for i, s in enumerate(symbols, 1):
        d.cell(row=i, column=30, value=s)             # AD 欄：商品清單（下拉選單來源）
    ws["A1"] = "跨週期最佳化 11 大指標參數庫"
    ws["A1"].font = TITLE
    ws["A2"] = "所有數值由程式從 MT5 K 線計算；參數只用前 70% 資料挑選。定義見「規則定義」分頁。"
    ws["A4"] = "選擇商品："
    ws["B5"] = dash_symbol
    ws["B5"].font = Font(bold=True, size=12, color="C00000")
    dv = DataValidation(type="list", formula1=f"='_dash'!$AD$1:$AD${len(symbols)}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B5")

    def block(top, header, prefix, ncols, first_col_is_symbol=True):
        for j, h in enumerate(header, 1):
            ws.cell(row=top, column=j, value=h)
        _header(ws, top, len(header))
        for k, tf in enumerate(TFS):
            rr = top + 1 + k
            if first_col_is_symbol:
                ws.cell(row=rr, column=1, value="=$B$5")
                ws.cell(row=rr, column=2, value=tf)
                start = 3
            else:
                ws.cell(row=rr, column=1, value=tf)
                start = 2
            key = f'"{prefix}|"&$B$5&"|"&"{tf}"'
            for j in range(ncols):
                col = get_column_letter(2 + j)          # _dash 的 v1 在 B 欄
                f = f'=IFERROR(INDEX(\'_dash\'!${col}:${col},MATCH({key},\'_dash\'!$A:$A,0)),"")'
                c = ws.cell(row=rr, column=start + j, value=f)
                if prefix == "S" and j == 0 or prefix == "P" and j == 18:
                    c.number_format = "yyyy-mm-dd hh:mm"

    block(6, ["商品 (Symbol)", "週期 (TF)", "最佳 SL", "最佳 TP", "短均線", "長均線", "RSI K數", "RSI下限", "RSI上限",
              "KD(K/D/平滑)", "PSY K數", "PSY上下限", "威廉 %R", "MTM K數", "MACD(快/慢/訊)", "布林(天/倍)", "CCI K數",
              "BIAS", "Keltner", "勝率/Alpha", "最後更新時間"], "P", 19)
    ws["A13"] = "【當前跨週期多空狀態與共振監控】"
    ws["A13"].font = Font(bold=True, size=12)
    block(14, ["商品 (Symbol)", "週期 (TF)", "最新 K線時間", "最新收盤價", "11指標多頭票", "11指標空頭票", "當前多空狀態",
               "週期權重", "跨週期共振訊號", "今日開盤價", "今日開盤至今漲跌%", "淨多空得分", "資料新鮮度", "盤勢 (ADX)",
               "分組綜合判定"], "S", 13)
    block(20, ["週期 (TF)", "1.MA", "2.RSI", "3.KD", "4.PSY", "5.威廉 %R", "6.MTM", "7.MACD", "8.布林", "9.CCI",
               "10.BIAS", "11.Keltner", "多頭總票", "空頭總票"], "V", 13, first_col_is_symbol=False)
    ws["A27"] = "註：『最佳 SL/TP』= S1（≥7/11 多數決）在前 70% 最佳的 ATR 倍數 × 目前 ATR(14)，單位為價格。"
    for col in range(1, 22):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["M"].width = 30
    ws.column_dimensions["O"].width = 18
    ws.freeze_panes = "A6"
    wb.active = 0
    wb.save(path)
