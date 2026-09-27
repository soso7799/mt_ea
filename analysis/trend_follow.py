"""
長週期趨勢跟隨（事先定好的單一設定，完整報告，不做參數搜尋）

  訊號：過去 63、126、252 個交易日（約 3、6、12 個月）報酬的正負號平均 → 部位方向 ∈ {-1,-1/3,+1/3,+1}
  調整：每週最後一個交易日收盤算訊號，下一個交易日開盤調整（以收盤報酬近似，訊號延後一天生效）
  部位：波動度目標。每商品權重 = 訊號 × (組合目標年化波動 10% ÷ √商品數) ÷ 該商品近 60 日年化波動
  成本：權重變動 × (點差% + 外匯手續費%)；未計隔夜利息
  對照：同樣風險配置但永遠做多
  穩健性（僅列出、不挑選）：只用 3 個月 / 6 個月 / 12 個月
  FTMO：最大回撤（限 10%）、單日最大虧損（限 5%），並算出要縮小到幾倍才能符合（留 20% 安全空間）

用法：python trend_follow.py --cache H:\\...\\export\\bars_mtf --out H:\\...\\reports [--offline]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from build_report import DEFAULT_SYMBOLS               # noqa: E402
from data import fetch_all                             # noqa: E402
from levels_study import is_fx                         # noqa: E402

LOOKBACKS = (63, 126, 252)
TARGET_VOL = 0.10
VOL_WIN = 60


def build_panel(frames, metas, syms, commission):
    close, cost = {}, {}
    for s in syms:
        df = frames.get((s, "D1"))
        if df is None or len(df) < 300:
            continue
        d = df.set_index(df["time"].dt.normalize())
        close[s] = d["close"]
        meta = metas.get(s, {})
        point = float(meta.get("point", 0.0001))
        spr = d["spread"].replace(0, np.nan).median()
        px = d["close"].median()
        c = (spr if not np.isnan(spr) else 1) * point / px
        if is_fx(s) and commission:
            contract = float(meta.get("contract") or 100000)
            # 每手手續費 ÷ 每手名目金額（美元）：XXXUSD = 合約 × 價格；USDXXX = 合約
            c += commission / (contract * px) if s.endswith("USD") else commission / contract
        cost[s] = c
    C = pd.DataFrame(close).sort_index()
    C = C[C.index.dayofweek < 5]
    return C, pd.Series(cost)


def signals(C, lookbacks):
    sig = sum(np.sign(C / C.shift(lb) - 1) for lb in lookbacks) / len(lookbacks)
    return sig


def simulate(C, cost, lookbacks, long_only=False):
    R = C.pct_change()
    n_assets = C.shape[1]
    vol = R.rolling(VOL_WIN, min_periods=40).std() * np.sqrt(252)
    raw = pd.DataFrame(1.0, index=C.index, columns=C.columns) if long_only else signals(C, lookbacks)
    w_target = raw * (TARGET_VOL / np.sqrt(n_assets)) / vol
    # 每週最後一個交易日決定，下一個交易日生效
    week = C.index.to_period("W")
    last_of_week = pd.Series(C.index, index=C.index).groupby(week).transform("max") == C.index
    W = w_target.where(last_of_week).ffill().shift(1).fillna(0.0)
    W = W.where(C.notna(), 0.0)
    gross = (W * R.fillna(0.0)).sum(axis=1)
    turnover = W.diff().abs().fillna(W.abs())
    costs = (turnover * cost.reindex(C.columns).fillna(0)).sum(axis=1)
    net = gross - costs
    contrib = (W * R.fillna(0.0)).sub(turnover * cost.reindex(C.columns).fillna(0))
    return net, W, contrib, costs


def stats(r):
    r = r.dropna()
    if len(r) < 20:
        return {}
    eq = (1 + r).cumprod()
    yrs = len(r) / 252
    dd = eq / eq.cummax() - 1
    m = (1 + r).groupby(r.index.to_period("M")).prod() - 1
    return dict(年化報酬=(eq.iloc[-1] ** (1 / yrs) - 1) * 100, 年化波動=r.std() * np.sqrt(252) * 100,
                夏普=r.mean() / r.std() * np.sqrt(252) if r.std() > 0 else np.nan,
                最大回撤=dd.min() * 100, 單日最大虧損=r.min() * 100, 最差月=m.min() * 100,
                月勝率=(m > 0).mean() * 100, 總報酬=(eq.iloc[-1] - 1) * 100, 年數=yrs)


def ftmo_scale(r, max_dd=0.10, max_day=0.05, margin=0.8):
    """要把部位乘上多少倍，才能讓最大回撤與單日虧損都在 FTMO 限制內（留 20% 空間）。"""
    lo, hi = 0.01, 5.0
    for _ in range(40):
        k = (lo + hi) / 2
        s = stats(r * k)
        ok = s and -s["最大回撤"] / 100 <= max_dd * margin and -s["單日最大虧損"] / 100 <= max_day * margin
        lo, hi = (k, hi) if ok else (lo, k)
    return lo


def main():
    ap = argparse.ArgumentParser(description="長週期趨勢跟隨（事先定好的單一設定）")
    ap.add_argument("--symbols", nargs="*", default=DEFAULT_SYMBOLS)
    ap.add_argument("--cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("--commission", type=float, default=5.0)
    args = ap.parse_args()

    frames, metas = fetch_all(args.symbols, args.cache, args.terminal, offline=args.offline, max_age_min=24 * 60)
    syms = [s for s in args.symbols if (s, "D1") in frames]
    C, cost = build_panel(frames, metas, syms, args.commission)
    start = C.index[252 + VOL_WIN]                      # 訊號與波動都有足夠資料之後才開始計算績效
    main_r, W, contrib, costs = simulate(C, cost, LOOKBACKS)
    main_r = main_r[main_r.index >= start]
    base_r, _, _, _ = simulate(C, cost, LOOKBACKS, long_only=True)
    base_r = base_r[base_r.index >= start]
    mid = main_r.index[len(main_r) // 2]

    rows = [dict(組合="趨勢跟隨（3/6/12 月平均）", 期間="全部", **stats(main_r)),
            dict(組合="趨勢跟隨（3/6/12 月平均）", 期間="前半", **stats(main_r[main_r.index < mid])),
            dict(組合="趨勢跟隨（3/6/12 月平均）", 期間="後半", **stats(main_r[main_r.index >= mid])),
            dict(組合="對照：永遠做多", 期間="全部", **stats(base_r)),
            dict(組合="對照：永遠做多", 期間="前半", **stats(base_r[base_r.index < mid])),
            dict(組合="對照：永遠做多", 期間="後半", **stats(base_r[base_r.index >= mid]))]
    for lb in LOOKBACKS:
        r_, _, _, _ = simulate(C, cost, (lb,))
        r_ = r_[r_.index >= start]
        rows.append(dict(組合=f"穩健性：只用 {lb // 21} 個月", 期間="全部", **stats(r_)))
    S = pd.DataFrame(rows)

    k = ftmo_scale(main_r)
    sk = stats(main_r * k)
    ftmo = pd.DataFrame([dict(項目="原始設定（目標年化波動 10%）", 倍數=1.0, **stats(main_r)),
                         dict(項目="縮放到符合 FTMO（回撤 ≤8%、單日 ≤4%）", 倍數=k, **sk)])

    c = contrib[contrib.index >= start]
    per = []
    for s in C.columns:
        x = c[s]
        per.append(dict(商品=s, 累積貢獻=x.sum() * 100, 夏普=x.mean() / x.std() * np.sqrt(252) if x.std() > 0 else np.nan,
                        前半貢獻=x[x.index < mid].sum() * 100, 後半貢獻=x[x.index >= mid].sum() * 100,
                        成本比例=cost[s] * 100))
    P = pd.DataFrame(per).sort_values("累積貢獻", ascending=False)
    yearly = pd.DataFrame({"趨勢跟隨%": ((1 + main_r).groupby(main_r.index.year).prod() - 1) * 100,
                           "永遠做多%": ((1 + base_r).groupby(base_r.index.year).prod() - 1) * 100})
    yearly.index.name = "年"
    sig_now = signals(C, LOOKBACKS).iloc[-1]
    vol_now = C.pct_change().rolling(VOL_WIN).std().iloc[-1] * np.sqrt(252)
    w_now = sig_now * (TARGET_VOL / np.sqrt(C.shape[1])) / vol_now
    now = pd.DataFrame({"方向": np.where(sig_now > 0, "做多", np.where(sig_now < 0, "做空", "空手")),
                        "訊號強度": sig_now.round(2), "年化波動%": (vol_now * 100).round(1),
                        "權重(名目/淨值)": w_now.round(3), "FTMO縮放後權重": (w_now * k).round(3),
                        "最新收盤日": C.apply(lambda s: s.last_valid_index())})
    now.index.name = "商品"

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"趨勢跟隨_{updated:%Y%m%d_%H%M}.xlsx")
    notes = [f"產生時間 {updated}；績效期間 {main_r.index.min().date()} ~ {main_r.index.max().date()}（前後半分界 {mid.date()}）；商品 {len(C.columns)} 個",
             "事先定好的單一設定，不做參數搜尋。",
             "訊號：過去 3、6、12 個月報酬正負號的平均（-1、-1/3、+1/3、+1）。每週最後一個交易日收盤決定，下一交易日起持有。",
             f"部位：波動度目標，每商品權重 = 訊號 × {TARGET_VOL:.0%}/√商品數 ÷ 近 {VOL_WIN} 日年化波動。權重 = 名目金額 ÷ 淨值。",
             "成本：權重變動量 × (點差% + 外匯手續費%)。未計隔夜利息（做多高利率貨幣有正利息，反之為負）。",
             "對照『永遠做多』：同樣的風險配置，只是永遠做多，用來分辨是趨勢判斷有用、還是只是市場上漲。",
             "穩健性：只用單一回看期的結果，僅供參考，不用來挑選。",
             "FTMO：原始設定若最大回撤 >10% 或單日虧損 >5%，另列縮小部位後的結果（留 20% 安全空間）。",
             "判定：趨勢跟隨的夏普在前半、後半都 >0.3，且全期高於『永遠做多』→ 趨勢判斷有額外價值。"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": notes}).to_excel(xw, sheet_name="說明", index=False)
        S.round(2).to_excel(xw, sheet_name="總結", index=False)
        ftmo.round(3).to_excel(xw, sheet_name="FTMO檢查", index=False)
        yearly.round(2).to_excel(xw, sheet_name="逐年報酬")
        P.round(3).to_excel(xw, sheet_name="各商品貢獻", index=False)
        now.to_excel(xw, sheet_name="目前部位")
        eq = pd.DataFrame({"趨勢跟隨": (1 + main_r).cumprod(), "永遠做多": (1 + base_r).cumprod()})
        eq.round(4).to_excel(xw, sheet_name="淨值曲線")
        from openpyxl.styles import Font, PatternFill
        for ws in xw.sheets.values():
            for cc in ws[1]:
                cc.font = Font(bold=True, color="FFFFFF")
                cc.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.5), 90)
            ws.freeze_panes = "A2"
    print(f"\n完成：{path}\n")
    for _, r_ in S.iterrows():
        print(f"  {r_['組合']:<22}{r_['期間']:<4} 年化 {r_.get('年化報酬', np.nan):+6.2f}%  波動 {r_.get('年化波動', np.nan):5.2f}%  "
              f"夏普 {r_.get('夏普', np.nan):+.2f}  最大回撤 {r_.get('最大回撤', np.nan):6.2f}%  單日最差 {r_.get('單日最大虧損', np.nan):5.2f}%")
    t_all, t1, t2 = stats(main_r)["夏普"], stats(main_r[main_r.index < mid])["夏普"], stats(main_r[main_r.index >= mid])["夏普"]
    ok = t1 > 0.3 and t2 > 0.3 and t_all > stats(base_r)["夏普"]
    print(f"\nFTMO：部位 × {k:.2f} 時，最大回撤 {sk['最大回撤']:.2f}%、單日最差 {sk['單日最大虧損']:.2f}%、年化 {sk['年化報酬']:+.2f}%")
    print(f"\n判定：{'✅ 趨勢判斷有額外價值' if ok else '❌ 不成立（前後半不穩定，或不如永遠做多）'}")


if __name__ == "__main__":
    main()
