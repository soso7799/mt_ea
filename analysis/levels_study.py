"""
關卡突破 / 反轉研究

關卡（每天用前一日資料算好、今天使用）：
  昨日 高/低/收；昨日亞盤、歐盤、美盤 高/低/收；上週 高/低/收
  （時段以 UTC 劃分：亞 00–07、歐 07–13、美 13–21；交易日 = MT5 伺服器日）
觸碰事件：價格從關卡一側（前一根收盤距關卡 >0.1 ATR）靠近並碰到關卡（±0.1 ATR），每條關卡每天只記第一次。
結果（觸碰 K 線收盤後，往後 N 根）：
  以下一根開盤（實際可進場價）為基準：突破 = 先往穿越方向走 b×ATR；反轉 = 先往回走 b×ATR；同一根兩者都碰到 = 不明；都沒有 = 未決
觸碰當下的判斷依據（只在接近關卡時才看指標）：
  成交量比（本根 tick volume ÷ 前 20 根平均）、觸碰 K 線是否收在關卡外、
  震盪指標同向極端（RSI>70 / K>80 / CCI>100 中 ≥2 個，於壓力處；支撐處反之）、
  ADX 盤勢、EMA50 趨勢方向、TAI 動能方向、關卡共振（0.3 ATR 內有其他關卡）、時段
交易規則：條件組合 × 突破/反轉 × SL/TP（ATR 倍數），前 70% 期間挑選、後 30% 驗證，扣點差與手續費。

用法：python levels_study.py --cache H:\\...\\export\\bars_mtf --out H:\\...\\reports [--tf M15] [--offline]
"""
import argparse
import datetime as dt
import itertools
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from analyze import closed_bars                         # noqa: E402
from build_report import DEFAULT_SYMBOLS                # noqa: E402
from data import TF_MIN, fetch_all                      # noqa: E402
from indicators import adx, atr, cci, ema, rsi, stoch, tai   # noqa: E402
from levels import server_to_utc                        # noqa: E402

SESS = [("亞", 0, 7), ("歐", 7, 13), ("美", 13, 21)]
LEVEL_NAMES = (["昨高", "昨低", "昨收"] +
               [f"昨{s}{k}" for s, _, _ in SESS for k in ("高", "低", "收")] +
               ["上週高", "上週低", "上週收"])
TOL, CONFL, BREAK_ATR = 0.1, 0.3, 0.5
SLTP = [(1.0, 1.5), (1.0, 2.0), (1.5, 2.0)]


def is_fx(sym):
    return len(sym) == 6 and sym.isalpha() and not sym.startswith(("XAU", "XAG"))


def daily_levels(df, utc_hour):
    """回傳每根 K 線所屬交易日的關卡（dict: name -> array）。"""
    day = df["time"].dt.normalize()
    week = day - pd.to_timedelta(day.dt.dayofweek, unit="D")
    g = df.assign(day=day, week=week, uh=utc_hour)
    D = g.groupby("day").agg(high=("high", "max"), low=("low", "min"), close=("close", "last"))
    for s, a, b in SESS:
        sub = g[(g["uh"] >= a) & (g["uh"] < b)].groupby("day").agg(
            **{f"{s}高": ("high", "max"), f"{s}低": ("low", "min"), f"{s}收": ("close", "last")})
        D = D.join(sub)
    P = D.shift(1)                                   # 前一個交易日
    W = g.groupby("week").agg(wh=("high", "max"), wl=("low", "min"), wc=("close", "last")).shift(1)
    lv = {"昨高": P["high"], "昨低": P["low"], "昨收": P["close"]}
    for s, _, _ in SESS:
        for k in ("高", "低", "收"):
            lv[f"昨{s}{k}"] = P[f"{s}{k}"]
    out = {k: v.reindex(day).to_numpy() for k, v in lv.items()}
    out["上週高"] = W["wh"].reindex(week).to_numpy()
    out["上週低"] = W["wl"].reindex(week).to_numpy()
    out["上週收"] = W["wc"].reindex(week).to_numpy()
    return out, day.to_numpy()


def find_events(sym, df, cost_price, N):
    o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
    tv = df["tick_volume"].to_numpy(float)
    n = len(c)
    utc = server_to_utc(df["time"])
    uh = utc.hour.to_numpy()
    lv, day = daily_levels(df, uh)
    a14 = atr(h, l, c, 14)
    r14 = rsi(c, 14)
    K, _ = stoch(h, l, c, 14, 3, 3)
    cc = cci(h, l, c, 20)
    ax, _, _ = adx(h, l, c, 14)
    e50 = ema(c, 50)
    col = tai(o, h, l, c, ma_period=28, tai_period=5)[1]
    volr = tv / pd.Series(tv).rolling(20, min_periods=20).mean().shift(1).to_numpy()
    L = np.vstack([lv[k] for k in LEVEL_NAMES])       # (15, n)
    ev = []
    day_start = np.r_[0, np.flatnonzero(day[1:] != day[:-1]) + 1]
    day_end = np.r_[day_start[1:], n]
    for k, name in enumerate(LEVEL_NAMES):
        for a0, b0 in zip(day_start, day_end):
            lvl = L[k, a0]
            if np.isnan(lvl):
                continue
            for i in range(max(a0, 1), b0):
                at = a14[i - 1]
                if np.isnan(at) or at <= 0:
                    continue
                tol = TOL * at
                if not (h[i] >= lvl - tol and l[i] <= lvl + tol):
                    continue
                if c[i - 1] < lvl - tol:
                    a = 1                      # 由下往上碰 → 壓力
                elif c[i - 1] > lvl + tol:
                    a = -1                     # 由上往下碰 → 支撐
                else:
                    continue                   # 前一根就在關卡附近，不算「靠近」，往後繼續找
                # 結果
                j0, j1 = i + 1, min(n, i + 1 + N)
                if j0 >= n:
                    break
                bk = BREAK_ATR * a14[i]
                ref = o[j0]                    # 從實際可進場價（下一根開盤）起算，避免「已收在關卡外」的量測偏差
                if a > 0:
                    brk, rev = h[j0:j1] >= ref + bk, l[j0:j1] <= ref - bk
                else:
                    brk, rev = l[j0:j1] <= ref - bk, h[j0:j1] >= ref + bk
                anyh = brk | rev
                if anyh.any():
                    x = int(np.argmax(anyh))
                    res = "不明" if brk[x] and rev[x] else ("突破" if brk[x] else "反轉")
                else:
                    res = "未決"
                ob = (r14[i] > 70) + (K[i] > 80) + (cc[i] > 100)
                os_ = (r14[i] < 30) + (K[i] < 20) + (cc[i] < -100)
                osc = "同向極端" if (ob >= 2 if a > 0 else os_ >= 2) else \
                      ("反向極端" if (os_ >= 2 if a > 0 else ob >= 2) else "非極端")
                slope = e50[i] - e50[i - 10] if i >= 10 else np.nan
                others = np.delete(L[:, i], k)
                dist = np.abs(others - lvl)
                confl = int(np.sum((dist <= CONFL * a14[i]) & (dist > 1e-9)))   # 價位完全相同的（如昨收=昨美收）不算共振
                sess = next((s for s, x0, x1 in SESS if x0 <= uh[i] < x1), "其他")
                ev.append(dict(
                    symbol=sym, time=df["time"].iloc[i], i=i, level=name, src=name[:-1], kind=name[-1],
                    price=lvl, a=a, side="壓力" if a > 0 else "支撐", result=res,
                    vol=volr[i], vol_b=("量縮<0.8" if volr[i] < 0.8 else ("量增>1.5" if volr[i] > 1.5 else
                                         ("略增1.2-1.5" if volr[i] > 1.2 else "正常"))) if not np.isnan(volr[i]) else "無資料",
                    close_beyond="收在關卡外" if (c[i] - lvl) * a > 0 else "收回關卡內",
                    osc=osc, adx=("趨勢盤" if ax[i] >= 25 else ("盤整盤" if ax[i] <= 20 else "不明確")),
                    trend=("推向關卡" if slope * a > 0 else "背離關卡") if not np.isnan(slope) else "無資料",
                    tai=("同向動能" if (col[i] == 1 and a > 0) or (col[i] == 2 and a < 0) else
                         ("反向動能" if col[i] in (1, 2) else "無動能")),
                    confl="共振" if confl else "單一", session=sess, atr=a14[i]))
                break                           # 每條關卡每天只記第一次觸碰
    E = pd.DataFrame(ev)
    if not len(E):
        return E
    # 交易結果（突破 / 反轉 × SL/TP），單位 R（1R = SL 距離）
    for mode in ("突破", "反轉"):
        for sl_m, tp_m in SLTP:
            key = f"{mode}_{sl_m}_{tp_m}"
            rs = []
            for _, e in E.iterrows():
                i, a = int(e["i"]), int(e["a"])
                d = a if mode == "突破" else -a
                if i + 1 >= n:
                    rs.append(np.nan)
                    continue
                px = o[i + 1]
                sl, tp = px - d * sl_m * e["atr"], px + d * tp_m * e["atr"]
                last = min(n - 1, i + 1 + 2 * N)
                hh, ll = h[i + 1:last + 1], l[i + 1:last + 1]
                s_hit, t_hit = (ll <= sl, hh >= tp) if d > 0 else (hh >= sl, ll <= tp)
                hit = s_hit | t_hit
                if hit.any():
                    x = int(np.argmax(hit))
                    exit_px = sl if s_hit[x] else tp
                else:
                    exit_px = c[last]
                rs.append(((exit_px - px) * d - cost_price(px)) / (sl_m * e["atr"]))
            E[key] = rs
    return E


def rate_table(E, by):
    rows = []
    for key, g in E.groupby(by, sort=False):
        n = len(g)
        b, r_ = (g["result"] == "突破").sum(), (g["result"] == "反轉").sum()
        dec = b + r_
        p = b / dec if dec else np.nan
        z = (b - dec * 0.5) / np.sqrt(dec * 0.25) if dec else np.nan
        rows.append(dict(**(dict(zip(by, key)) if isinstance(key, tuple) else {by[0]: key}),
                         事件數=n, 突破=b, 反轉=r_, 未決或不明=n - dec, 突破率=p * 100 if dec else np.nan,
                         反轉率=(1 - p) * 100 if dec else np.nan, z值=z,
                         傾向=("偏突破" if z >= 2 else ("偏反轉" if z <= -2 else "無明顯傾向")) if dec else ""))
    return pd.DataFrame(rows)


FILTERS = {
    "level": {"全部": None, "只含高低": lambda E: E["kind"] != "收", "只含收盤": lambda E: E["kind"] == "收",
              "上週": lambda E: E["src"] == "上週", "昨日全日": lambda E: E["src"] == "昨",
              "昨日時段": lambda E: E["src"].isin(["昨亞", "昨歐", "昨美"])},
    "vol": {"不限": None, "量縮<0.8": lambda E: E["vol"] < 0.8, "量增>1.2": lambda E: E["vol"] > 1.2,
            "量增>1.5": lambda E: E["vol"] > 1.5},
    "close": {"不限": None, "收在關卡外": lambda E: E["close_beyond"] == "收在關卡外",
              "收回關卡內": lambda E: E["close_beyond"] == "收回關卡內"},
    "osc": {"不限": None, "同向極端": lambda E: E["osc"] == "同向極端", "非極端": lambda E: E["osc"] != "同向極端"},
    "adx": {"不限": None, "趨勢盤": lambda E: E["adx"] == "趨勢盤", "盤整盤": lambda E: E["adx"] == "盤整盤"},
    "confl": {"不限": None, "共振": lambda E: E["confl"] == "共振"},
    "tai": {"不限": None, "同向動能": lambda E: E["tai"] == "同向動能", "非同向": lambda E: E["tai"] != "同向動能"},
}


def tstat(x):
    x = x[~np.isnan(x)]
    if len(x) < 2 or x.std(ddof=1) == 0:
        return len(x), (x.mean() if len(x) else np.nan), np.nan, np.nan
    return len(x), x.mean(), x.mean() / x.std(ddof=1) * np.sqrt(len(x)), (x > 0).mean() * 100


def search_rules(E, cut):
    masks = {}
    for fname, opts in FILTERS.items():
        masks[fname] = {k: (np.ones(len(E), bool) if f is None else f(E).to_numpy()) for k, f in opts.items()}
    is_m = (E["time"] < cut).to_numpy()
    out = []
    names = list(FILTERS)
    seen = set()
    sym = E["symbol"].to_numpy()
    for combo in itertools.product(*[list(FILTERS[f]) for f in names]):
        m = np.ones(len(E), bool)
        for f, v in zip(names, combo):
            m &= masks[f][v]
        if m.sum() < 60:
            continue
        h = hash(np.packbits(m).tobytes())      # 條件不同但實際選到同一批事件 → 視為同一條
        if h in seen:
            continue
        seen.add(h)
        for mode in ("突破", "反轉"):
            for sl_m, tp_m in SLTP:
                r = E[f"{mode}_{sl_m}_{tp_m}"].to_numpy()
                n1, m1, t1, w1 = tstat(r[m & is_m])
                if n1 < 50 or not (m1 > 0):
                    continue
                oo = m & ~is_m
                n2, m2, t2, w2 = tstat(r[oo])
                pos, cnt = 0, 0
                for s_ in np.unique(sym[oo]):
                    x = r[oo & (sym == s_)]
                    x = x[~np.isnan(x)]
                    if len(x) >= 5:
                        cnt += 1
                        pos += x.mean() > 0
                out.append(dict(zip(["關卡", "成交量", "觸碰K收盤", "震盪指標", "盤勢", "共振", "TAI"], combo),
                                做法=mode, SL=sl_m, TP=tp_m, 前70筆數=n1, 前70平均R=m1, 前70t=t1, 前70勝率=w1,
                                後30筆數=n2, 後30平均R=m2, 後30t=t2, 後30勝率=w2,
                                後30獲利商品=f"{pos}/{cnt}",
                                可信=bool(t1 >= 2 and n2 >= 50 and t2 >= 3 and cnt >= 3 and pos / cnt >= 0.6)))
    R = pd.DataFrame(out)
    if len(R):
        R = R.sort_values("前70t", ascending=False)
    return R


def today_levels(sym, df, E, digits):
    """今天（最後一個交易日）的關卡清單、距現價、歷史突破/反轉率。"""
    o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
    utc = server_to_utc(df["time"])
    lv, _ = daily_levels(df, utc.hour.to_numpy())
    a14 = atr(h, l, c, 14)[-1]
    px = c[-1]
    base = E.groupby("level")["result"].agg(lambda s: ((s == "突破").sum(), (s == "反轉").sum()))
    rows = []
    for name in LEVEL_NAMES:
        v = lv[name][-1]
        if np.isnan(v):
            continue
        others = [lv[x][-1] for x in LEVEL_NAMES if x != name and not np.isnan(lv[x][-1])]
        cf = sum(1e-9 < abs(x - v) <= CONFL * a14 for x in others)
        b, r_ = base.get(name, (0, 0))
        rows.append(dict(商品=sym, 關卡=name, 價位=round(v, digits), 現價=round(px, digits),
                         距離ATR=round((v - px) / a14, 2), 位置="上方壓力" if v > px else "下方支撐",
                         共振關卡數=cf, 此關卡歷史突破率=round(b / (b + r_) * 100, 1) if b + r_ else None,
                         歷史樣本=b + r_))
    return sorted(rows, key=lambda r: abs(r["距離ATR"]))


def main():
    ap = argparse.ArgumentParser(description="關卡突破 / 反轉研究")
    ap.add_argument("--symbols", nargs="*", default=DEFAULT_SYMBOLS)
    ap.add_argument("--cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("--tf", default="M15", choices=["M5", "M15", "H1"], help="用哪個週期偵測觸碰（預設 M15）")
    ap.add_argument("--hours", type=float, default=3, help="觸碰後觀察幾小時判斷突破/反轉（預設 3）")
    ap.add_argument("--commission", type=float, default=5.0, help="外匯每手來回手續費 USD（FTMO 預設 5）")
    args = ap.parse_args()

    frames, metas = fetch_all(args.symbols, args.cache, args.terminal, offline=args.offline, max_age_min=24 * 60)
    syms = [s for s in args.symbols if (s, args.tf) in frames]
    ticks = [pd.Timestamp(m["tick_time"]) for m in metas.values() if m.get("tick_time")]
    server_now = max(ticks) if ticks else None
    N = max(1, int(round(args.hours * 60 / TF_MIN[args.tf])))
    allE, today = [], []
    for s in syms:
        meta = metas.get(s, {})
        df = closed_bars(frames[(s, args.tf)], args.tf, server_now)
        point = float(meta.get("point", 0.0001))
        spr = df["spread"].replace(0, np.nan).median()
        spread_price = (spr if not np.isnan(spr) else 1) * point
        contract = float(meta.get("contract") or 100000)
        if is_fx(s) and args.commission:
            if s.endswith("USD"):
                cost = lambda px, sp=spread_price, cm=args.commission / contract: sp + cm          # noqa: E731
            else:
                cost = lambda px, sp=spread_price, cm=args.commission / contract: sp + cm * px     # noqa: E731
        else:
            cost = lambda px, sp=spread_price: sp                                                # noqa: E731
        E = find_events(s, df, cost, N)
        print(f"  {s}: {len(E)} 個觸碰事件")
        if len(E):
            allE.append(E)
            today += today_levels(s, df, E, int(meta.get("digits", 5)))
    E = pd.concat(allE, ignore_index=True).sort_values("time").reset_index(drop=True)
    cut = E["time"].quantile(0.70)

    base = rate_table(E, ["level"])
    cond = pd.concat([rate_table(E, [f]).rename(columns={f: "數值"}).assign(條件=name)
                      for f, name in [("vol_b", "成交量比"), ("close_beyond", "觸碰K收盤"), ("osc", "震盪指標"),
                                      ("adx", "ADX盤勢"), ("trend", "EMA50趨勢"), ("tai", "TAI動能"),
                                      ("confl", "關卡共振"), ("session", "觸碰時段"), ("side", "壓力或支撐")]],
                     ignore_index=True)
    cond = cond[["條件", "數值"] + [c for c in cond.columns if c not in ("條件", "數值")]]
    combo = rate_table(E, ["vol_b", "close_beyond"])
    R = search_rules(E, cut)
    sym_best = []
    if len(R):
        top = R.iloc[0]
        for s, g in E.groupby("symbol"):
            m = np.ones(len(g), bool)
            for f, col in zip(FILTERS, ["關卡", "成交量", "觸碰K收盤", "震盪指標", "盤勢", "共振", "TAI"]):
                fn = FILTERS[f][top[col]]
                if fn is not None:
                    m &= fn(g).to_numpy()
            r = g[f"{top['做法']}_{top['SL']}_{top['TP']}"].to_numpy()[m]
            is_m = (g["time"] < cut).to_numpy()[m]
            n1, m1, t1, _ = tstat(r[is_m])
            n2, m2, t2, _ = tstat(r[~is_m])
            sym_best.append(dict(商品=s, 前70筆數=n1, 前70平均R=m1, 後30筆數=n2, 後30平均R=m2, 後30t=t2))

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"關卡突破反轉_{args.tf}_{updated:%Y%m%d_%H%M}.xlsx")
    notes = [
        f"產生時間 {updated}；週期 {args.tf}；觸碰後觀察 {args.hours} 小時（{N} 根）；商品：{', '.join(syms)}",
        f"期間 {E['time'].min()} ~ {E['time'].max()}；前 70% / 後 30% 分界 {cut}",
        "關卡：昨日高/低/收；昨日亞盤(UTC00-07)、歐盤(07-13)、美盤(13-21)的高/低/收；上週高/低/收。每天用前一日資料算好後使用。",
        f"觸碰：前一根收盤距關卡 >{TOL} ATR，本根高低點碰到關卡 ±{TOL} ATR；每條關卡每天只記第一次。",
        f"突破 / 反轉：以觸碰 K 線的下一根開盤為基準，先往穿越方向走 {BREAK_ATR} ATR = 突破，先往回走 {BREAK_ATR} ATR = 反轉。"
        "z 值 ≥2 偏突破、≤-2 偏反轉（相對 50/50）。",
        "成交量比 = 觸碰 K 線 tick volume ÷ 前 20 根平均。震盪同向極端 = 碰壓力時 RSI>70/K>80/CCI>100 有 ≥2 個（碰支撐反之）。",
        f"共振 = {CONFL} ATR 內還有其他關卡。EMA50 推向關卡 = 近 10 根 EMA50 斜率朝關卡方向。",
        "交易：觸碰 K 線收盤後下一根開盤進場；突破 = 順著穿越方向做，反轉 = 反方向做；SL/TP 為 ATR 倍數；"
        f"最多持有 {2 * N} 根。報酬單位 R（1R = 停損距離），已扣點差" + (f"與外匯手續費 {args.commission} USD/手。" if args.commission else "。"),
        "規則搜尋：所有條件組合 × 突破/反轉 × SL/TP，只用前 70% 挑選（筆數≥50、平均R>0，依 t 值排序），後 30% 為樣本外驗證。",
        "注意：組合數上千，前 70% 最好的結果必然偏樂觀；後 30% 以 t≥2 篩選時，純運氣約有 2.3% 的規則會通過。",
        "可信 = 前 70% t≥2，且後 30% 筆數≥50、t≥3、≥3 個商品中 ≥60% 為正。兩段都顯著才算，只有『可信』的規則值得寫成 EA。",
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": notes}).to_excel(xw, sheet_name="說明", index=False)
        pd.DataFrame(today).to_excel(xw, sheet_name="今日關卡", index=False)
        base.round(2).to_excel(xw, sheet_name="各關卡突破反轉率", index=False)
        cond.round(2).to_excel(xw, sheet_name="觸碰條件分析", index=False)
        combo.round(2).rename(columns={"vol_b": "成交量比", "close_beyond": "觸碰K收盤"}).to_excel(
            xw, sheet_name="成交量x收盤", index=False)
        (R.head(300).round(3) if len(R) else pd.DataFrame({"結果": ["沒有規則在前 70% 達標"]})).to_excel(
            xw, sheet_name="規則回測", index=False)
        pd.DataFrame(sym_best).round(3).to_excel(xw, sheet_name="最佳規則各商品", index=False)
        E.drop(columns=["i"]).tail(5000).round({"price": 6, "vol": 3, "atr": 6}).to_excel(xw, sheet_name="事件明細(最近5000)", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.sheets.values():
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.5), 70)
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.value == "偏突破":
                        cell.fill = PatternFill("solid", fgColor="C6EFCE")
                    elif cell.value == "偏反轉":
                        cell.fill = PatternFill("solid", fgColor="BDD7EE")
    print(f"\n完成：{path}")
    print(f"\n觸碰事件共 {len(E)} 個；整體 突破 {(E['result'] == '突破').mean():.1%}、反轉 {(E['result'] == '反轉').mean():.1%}")
    print("\n各條件的突破/反轉傾向（|z|≥2）：")
    for _, r_ in cond[cond["傾向"].isin(["偏突破", "偏反轉"])].iterrows():
        print(f"  {r_['條件']}={r_['數值']:<10} 突破率 {r_['突破率']:.1f}%  z={r_['z值']:+.1f}  n={r_['突破'] + r_['反轉']}")
    if len(R):
        ok2 = R[(R["後30平均R"] > 0) & (R["後30t"] >= 2)]
        ok = R[R["可信"]]
        print(f"\n規則：前 70% 達標 {len(R)} 條（已合併重複）；後 30% 也為正且 t≥2 的有 {len(ok2)} 條，"
              f"但純靠運氣預期就有約 {len(R) * 0.023:.0f} 條。")
        print(f"可信（前 70% t≥2，且後 30% t≥3、≥50 筆、≥60% 商品獲利）：{len(ok)} 條" + ("。前 10：" if len(ok) else "。"))
        for _, r_ in ok.head(10).iterrows():
            print(f"  {r_['做法']} | 關卡:{r_['關卡']} 量:{r_['成交量']} 收盤:{r_['觸碰K收盤']} 震盪:{r_['震盪指標']} "
                  f"盤:{r_['盤勢']} 共振:{r_['共振']} TAI:{r_['TAI']} SL{r_['SL']}/TP{r_['TP']} | "
                  f"前70 {r_['前70平均R']:+.3f}R(n={r_['前70筆數']}) 後30 {r_['後30平均R']:+.3f}R t={r_['後30t']:.1f}(n={r_['後30筆數']})")


if __name__ == "__main__":
    main()
