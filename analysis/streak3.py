"""
連續 3 根 K 不反轉就做 —— 回測（所有商品、M30）

訊號（只用已收盤 K 線）：最近 3 根都是陽線（收 > 開）且再前一根不是陽線 → 下一根開盤做多；陰線反之做空
事先定好的出場（參數不優化）：
  A 出現反轉K就出：持有到第一根反向 K 線（陽線單遇到陰線）收盤，下一根開盤出場；安全停損 2 ATR；最多 48 根
  B 固定停損停利：SL 1 ATR、TP 2 ATR，最多 48 根
對照：同樣訊號反向做（賭 3 根後會回檔）
另列 4 根、5 根連續 K 的結果（只供參考）
成交：下一根開盤、Bid/Ask、扣點差與外匯手續費；同商品一次一筆；同根碰到 SL/TP 算 SL
FTMO 試算：每筆風險 0.5% 帳戶，所有商品合計

用法：python streak3.py --m1-cache H:\\...\\export\\bars --out H:\\...\\reports [--years 2] [--tf M30]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "ml"))

from indicators import atr                    # noqa: E402
from levels_study import is_fx                # noqa: E402
from mine_all import discover_symbols         # noqa: E402
import optimize_wf as ow                      # noqa: E402

MAX_HOLD = 48
RISK_PCT = 0.5
MINUTES = {"M15": 15, "M30": 30, "H1": 60}


def streak_signals(o, c, k):
    bull = c > o
    bear = c < o
    bu = pd.Series(bull.astype(float)).rolling(k).sum().to_numpy() == k
    be = pd.Series(bear.astype(float)).rolling(k).sum().to_numpy() == k
    return ow.onset(bu), ow.onset(be)


def sim(entries, o, h, l, c, sp, A, comm_px, mode):
    """mode 'A'：反轉K出場 + 2ATR 安全停損；'B'：SL1/TP2。回傳 list of (e, x, d, R, R0)。"""
    n = len(c)
    out, busy = [], -1
    for i, d in entries:
        e = i + 1
        if e <= busy or e >= n - 1 or not A[i] > 0:
            continue
        entry = o[e] + (sp[e] if d > 0 else 0.0)
        sl_m = 2.0 if mode == "A" else 1.0
        sl = entry - d * sl_m * A[i]
        tp = entry + d * 2.0 * A[i] if mode == "B" else np.nan
        dist = sl_m * A[i]
        j = np.arange(e, min(n, e + MAX_HOLD))
        adj = 0.0 if d > 0 else sp[j]
        hi, lo = h[j] + adj, l[j] + adj
        s_hit = (lo <= sl) if d > 0 else (hi >= sl)
        t_hit = np.zeros(len(j), bool) if np.isnan(tp) else ((hi >= tp) if d > 0 else (lo <= tp))
        rev = ((c[j] - o[j]) * d < 0) if mode == "A" else np.zeros(len(j), bool)
        ev = s_hit | t_hit | rev
        if ev.any():
            k = int(np.argmax(ev))
            x = j[k]
            if s_hit[k]:
                op = o[x] + (0.0 if d > 0 else sp[x])
                px = op if (op - sl) * d < 0 else sl
            elif t_hit[k]:
                px = tp
            else:                                   # 反轉 K 收盤 → 下一根開盤出
                x = min(x + 1, n - 1)
                px = o[x] + (0.0 if d > 0 else sp[x])
        else:
            x = j[-1]
            px = c[x] + (0.0 if d > 0 else sp[x])
        R = ((px - entry) * d - comm_px) / dist
        R0 = R + (comm_px + (sp[e] if d > 0 else sp[x])) / dist
        out.append((e, x, d, R, R0))
        busy = x
    return out


def stat(R):
    R = np.asarray(R, float)
    n = len(R)
    if n < 2:
        return n, (R.mean() if n else np.nan), np.nan
    sd = R.std(ddof=1)
    return n, R.mean(), (R.mean() / sd * np.sqrt(n) if sd > 1e-9 else np.nan)


def main():
    ap = argparse.ArgumentParser(description="連續 3 根 K 不反轉就做 —— 回測")
    ap.add_argument("--symbols", nargs="*", default=None, help="預設 = 快取 + MT5 市場報價顯示的所有商品")
    ap.add_argument("--m1-cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--years", type=float, default=2.0)
    ap.add_argument("--to", default=None)
    ap.add_argument("--tf", default="M30", choices=list(MINUTES))
    ap.add_argument("--commission", type=float, default=5.0)
    args = ap.parse_args()

    from mt5data import fetch_m1
    end = dt.datetime.fromisoformat(args.to) if args.to else dt.datetime.combine(dt.date.today(), dt.time())
    start = end - dt.timedelta(days=round(args.years * 365.25))
    syms = args.symbols or discover_symbols(args.m1_cache)
    print(f"商品 {len(syms)} 個；{args.tf}；{start.date()} ~ {end.date()}")
    rows = []
    for s in syms:
        try:
            got = fetch_m1([s], start - dt.timedelta(days=10), end, args.m1_cache, args.terminal)
        except SystemExit as e:
            print(f"⚠️ {s} 略過：{e}")
            continue
        m1, meta = got[s]
        m1 = m1[(m1["time"] >= pd.Timestamp(start)) & (m1["time"] < pd.Timestamp(end))]
        if len(m1) < 50000:
            continue
        df = ow.resample(m1, MINUTES[args.tf])
        o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
        t = df["time"].to_numpy()
        point = float(meta.get("point", 0.0001))
        contract = float(meta.get("contract") or 100000)
        sp = df["spread"].to_numpy(float) * point
        med = np.nanmedian(sp[sp > 0]) if (sp > 0).any() else point
        sp = np.where(sp > 0, sp, med)
        comm_px = args.commission / contract * (1.0 if s.endswith("USD") else np.nanmedian(c)) if is_fx(s) and args.commission else 0.0
        A = atr(h, l, c, 14)
        for k in (3, 4, 5):
            bu, be = streak_signals(o, c, k)
            ent = sorted([(i, 1) for i in np.flatnonzero(bu)] + [(i, -1) for i in np.flatnonzero(be)])
            for mode, label in (("A", "A 反轉K就出"), ("B", "B SL1/TP2")):
                for dirn, dname in ((1, "順勢做"), (-1, "反向做(對照)")):
                    for e, x, d, R, R0 in sim([(i, d * dirn) for i, d in ent], o, h, l, c, sp, A, comm_px, mode):
                        rows.append(dict(商品=s, 連續根數=k, 出場=label, 方向=dname, 進場時間=t[e], 出場時間=t[x],
                                         多空="多" if d > 0 else "空", R=R, 不含成本R=R0, 持有根數=x - e))
        print(f"  {s}: 完成")
    if not rows:
        raise SystemExit("沒有資料")
    T = pd.DataFrame(rows)
    mid = T["進場時間"].quantile(0.5)
    summ, per = [], []
    for (k, ex, dn), g in T.groupby(["連續根數", "出場", "方向"]):
        n, mu, tt = stat(g["R"])
        h1 = g.loc[g["進場時間"] < mid, "R"].mean()
        h2 = g.loc[g["進場時間"] >= mid, "R"].mean()
        bys = g.groupby("商品")["R"].sum()
        s_ = g.sort_values("出場時間")
        eq = np.cumsum(s_["R"].to_numpy() * RISK_PCT)
        dd = (np.maximum.accumulate(np.r_[0, eq])[1:] - eq).max()
        worst = (s_.groupby(pd.to_datetime(s_["出場時間"]).dt.date)["R"].sum() * RISK_PCT).min()
        ok = k == 3 and not dn.startswith("反向") and n >= 100 and mu > 0 and (tt or 0) >= 2 and h1 > 0 and h2 > 0 \
            and (bys > 0).mean() >= 0.6
        summ.append(dict(連續根數=k, 出場=ex, 方向=dn, 筆數=n, 平均R=mu, t值=tt, 不含成本平均R=g["不含成本R"].mean(),
                         勝率=(g["R"] > 0).mean() * 100, 前半平均R=h1, 後半平均R=h2, 獲利商品比例=(bys > 0).mean() * 100,
                         平均持有根數=g["持有根數"].mean(), **{"帳戶報酬%(每筆0.5%)": eq[-1] if len(eq) else 0,
                                                     "最大回撤%": dd, "最差單日%": worst},
                         判定=("✅ 可用" if ok else ("（參考）" if k != 3 else ("對照" if dn.startswith("反向") else "❌")))))
    S = pd.DataFrame(summ)
    for (s, ex, dn), g in T[T["連續根數"] == 3].groupby(["商品", "出場", "方向"]):
        n, mu, tt = stat(g["R"])
        per.append(dict(商品=s, 出場=ex, 方向=dn, 筆數=n, 平均R=mu, t值=tt, 不含成本平均R=g["不含成本R"].mean(),
                        前半平均R=g.loc[g["進場時間"] < mid, "R"].mean(), 後半平均R=g.loc[g["進場時間"] >= mid, "R"].mean(),
                        總R=g["R"].sum()))
    Pm = pd.DataFrame(per)

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"連續K回測_{args.tf}_{updated:%Y%m%d_%H%M}.xlsx")
    notes = [f"產生時間 {updated}；{args.tf}；期間 {start.date()} ~ {end.date()}；商品 {T['商品'].nunique()} 個",
             "訊號：最近 3 根都是陽線（且再前一根不是）→ 下一根開盤做多；3 根陰線 → 做空。只用已收盤 K 線。",
             "A 反轉K就出：持有到第一根反向 K 收盤、下一根開盤出場；安全停損 2 ATR；最多 48 根。B：SL 1 ATR / TP 2 ATR。",
             "對照：同樣訊號反向做。4 根、5 根連續 K 只供參考（沒有事先指定，不做判定）。",
             "成本：點差 + 外匯手續費；『不含成本平均R』看規則本身。1R = 停損距離。",
             "判定 ✅：3 根、順勢、≥100 筆、平均R>0、t≥2、前後半都賺、≥60% 商品獲利。",
             "『各商品』只供參考：單一商品挑最好的很容易是運氣，要看前半 / 後半是否都賺。",
             f"監視程式（monitor_streak.bat）會讀這個檔案的『各商品』，在訊號旁標出該商品的回測平均R。"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": notes}).to_excel(xw, sheet_name="說明", index=False)
        S.round(3).to_excel(xw, sheet_name="總結", index=False)
        Pm.round(3).to_excel(xw, sheet_name="各商品", index=False)
        T.tail(20000).round(3).to_excel(xw, sheet_name="交易明細", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.sheets.values():
            for cc in ws[1]:
                cc.font = Font(bold=True, color="FFFFFF")
                cc.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.3), 50)
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("✅"):
                        cell.fill = PatternFill("solid", fgColor="C6EFCE")
    Pm.to_csv(os.path.join(args.out, f"連續K回測_{args.tf}_各商品_latest.csv"), index=False, encoding="utf-8-sig")
    print(f"\n完成：{path}\n")
    for _, r in S[S["連續根數"] == 3].iterrows():
        print(f"  3根 {r['出場']:<9} {r['方向']:<8} {r['筆數']:>6} 筆  平均 {r['平均R']:+.3f}R（不含成本 {r['不含成本平均R']:+.3f}R）"
              f"  t={r['t值']:.2f}  前半 {r['前半平均R']:+.3f} 後半 {r['後半平均R']:+.3f}  獲利商品 {r['獲利商品比例']:.0f}%  → {r['判定']}")


if __name__ == "__main__":
    main()
