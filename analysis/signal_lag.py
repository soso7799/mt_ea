"""
指標訊號落後多少？—— 訊號出現時，行情已經走了多少、還剩多少

對你圖上用的每個訊號（你目前的設定），在訊號 K 線收盤時量：
  落後根數：距離「最近 48 根內真正的轉折點」（做多 = 最低點、做空 = 最高點）已經幾根
  已走幅度：從轉折點到訊號收盤已經走了幾個 ATR
  之後最大順向 / 最大逆向：下一根開盤進場後 48 根內，最多往有利 / 不利方向走幾個 ATR
  已走比例：已走 ÷（已走 + 之後最大順向）—— 訊號出現時，整段行情已經用掉多少
  之後 6 / 12 / 48 根報酬（ATR）：不含成本
對照：隨機時間、隨機方向進場的同樣數字。訊號若真的「提早」，已走幅度應小於隨機、之後順向應大於逆向。
全部只用訊號 K 線收盤前的資料（沒有偷看未來），和 MT5 圖上已收盤 K 線的箭頭相同。

用法：python signal_lag.py --m1-cache H:\\...\\export\\bars --out H:\\...\\reports [--years 1] [--tf M5]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "ml"))

from build_report import DEFAULT_SYMBOLS      # noqa: E402
from indicators import atr                    # noqa: E402
import optimize_wf as ow                      # noqa: E402

W = 48
# 你目前的設定（觸發名稱, 參數）
YOURS = [("MA交叉", "EMA20/50"), ("MACD Hull", "5/35/5"), ("MACD柱", "12/26 穿0"), ("MACD柱", "12/26 柱轉折"),
         ("KDJ J線", "9/3/3 J=3D-2K 穿50"), ("KDJ J線", "9/3/3 J=3K-2D 穿50"), ("KDJ", "9/3/3 區間20/80"),
         ("TAI", "MA28/週期5"), ("趨勢線突破", "Swing3")]


def measure(idx, d, o, h, l, c, A):
    n = len(c)
    ok = (idx >= W) & (idx + W + 1 < n) & (A[idx] > 0)
    idx, d = idx[ok], d[ok]
    if not len(idx):
        return pd.DataFrame()
    back = idx[:, None] - np.arange(W)[None, :]            # i, i-1, ..., i-W+1
    fwd = idx[:, None] + 1 + np.arange(W)[None, :]         # i+1 ... i+W
    a = A[idx]
    lo_b, hi_b = l[back], h[back]
    k_lo, k_hi = np.argmin(lo_b, 1), np.argmax(hi_b, 1)
    long_ = d > 0
    lag = np.where(long_, k_lo, k_hi)
    done = np.where(long_, c[idx] - lo_b.min(1), hi_b.max(1) - c[idx]) / a
    entry = o[idx + 1]
    mfe = np.where(long_, h[fwd].max(1) - entry, entry - l[fwd].min(1)) / a
    mae = np.where(long_, entry - l[fwd].min(1), h[fwd].max(1) - entry) / a
    out = dict(落後根數=lag, 已走幅度ATR=done, 之後最大順向ATR=mfe, 之後最大逆向ATR=mae,
               已走比例=done / np.where(done + mfe > 0, done + mfe, np.nan))
    for N in (6, 12, 48):
        out[f"之後{N}根報酬ATR"] = (c[idx + N] - entry) * d / a
    return pd.DataFrame(out)


def main():
    ap = argparse.ArgumentParser(description="指標訊號落後多少")
    ap.add_argument("--symbols", nargs="*", default=DEFAULT_SYMBOLS)
    ap.add_argument("--m1-cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--years", type=float, default=1.0)
    ap.add_argument("--to", default=None)
    ap.add_argument("--tf", default="M5", choices=["M5", "M15", "H1"])
    args = ap.parse_args()

    from mt5data import fetch_m1
    end = dt.datetime.fromisoformat(args.to) if args.to else dt.datetime.combine(dt.date.today(), dt.time())
    start = end - dt.timedelta(days=round(args.years * 365.25))
    minutes = {"M5": 5, "M15": 15, "H1": 60}[args.tf]
    rng = np.random.default_rng(0)
    rows = []
    for s in args.symbols:
        try:
            got = fetch_m1([s], start - dt.timedelta(days=10), end, args.m1_cache, args.terminal)
        except SystemExit as e:
            print(f"⚠️ {s} 略過：{e}")
            continue
        m1, _ = got[s]
        m1 = m1[(m1["time"] >= pd.Timestamp(start)) & (m1["time"] < pd.Timestamp(end))]
        df = ow.resample(m1, minutes)
        if len(df) < 2000:
            continue
        o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
        A = atr(h, l, c, 14)
        want = set(YOURS)
        for trig, prm, lg, sh in ow.triggers(o, h, l, c, df["time"].to_numpy()):
            if (trig, prm) not in want:
                continue
            L, S = np.flatnonzero(lg), np.flatnonzero(sh)
            M = measure(np.r_[L, S], np.r_[np.ones(len(L)), -np.ones(len(S))], o, h, l, c, A)
            if len(M):
                rows.append(M.assign(商品=s, 訊號=f"{trig} {prm}"))
        k = rng.integers(W, len(c) - W - 2, size=5000)
        M = measure(k, rng.choice([-1.0, 1.0], size=len(k)), o, h, l, c, A)
        rows.append(M.assign(商品=s, 訊號="（對照）隨機進場"))
        print(f"  {s}: 完成")
    if not rows:
        raise SystemExit("沒有資料")
    D = pd.concat(rows, ignore_index=True)
    cols = ["落後根數", "已走幅度ATR", "之後最大順向ATR", "之後最大逆向ATR", "已走比例", "之後6根報酬ATR", "之後12根報酬ATR", "之後48根報酬ATR"]
    S = D.groupby("訊號")[cols].median().add_prefix("中位數_")
    S.insert(0, "筆數", D.groupby("訊號").size())
    S["平均_之後12根報酬ATR"] = D.groupby("訊號")["之後12根報酬ATR"].mean()
    S["平均_之後48根報酬ATR"] = D.groupby("訊號")["之後48根報酬ATR"].mean()
    S["順向>逆向比例"] = D.assign(x=D["之後最大順向ATR"] > D["之後最大逆向ATR"]).groupby("訊號")["x"].mean() * 100
    S = S.reset_index()
    base = S[S["訊號"].str.startswith("（對照）")]
    P = D.groupby(["商品", "訊號"])[cols].median().reset_index()

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"訊號落後分析_{updated:%Y%m%d_%H%M}.xlsx")
    notes = [f"產生時間 {updated}；期間 {start.date()} ~ {end.date()}；{args.tf}；ATR(14) 為單位",
             "落後根數：訊號收盤時，距離最近 48 根內的真正轉折點（做多看最低點、做空看最高點）已經幾根。",
             "已走幅度：轉折點到訊號收盤已經走了幾個 ATR；已走比例 = 已走 ÷（已走 + 之後最大順向）。",
             "之後最大順向 / 逆向：下一根開盤進場後 48 根內最有利 / 最不利走多遠；順向>逆向比例 > 50% 才表示訊號後行情還偏向你。",
             "之後 N 根報酬：不含成本；平均接近 0 = 訊號出現時，後面漲跌已經跟隨機一樣。",
             "（對照）隨機進場：隨機時間、隨機方向。訊號若真的提早，『已走幅度』應明顯小於對照、之後順向應大於逆向。",
             "只用訊號 K 收盤前的資料，同 MT5 圖上已收盤 K 線的箭頭（這些指標都不重繪）。"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": notes}).to_excel(xw, sheet_name="說明", index=False)
        S.round(3).to_excel(xw, sheet_name="總結", index=False)
        P.round(3).to_excel(xw, sheet_name="各商品", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.sheets.values():
            for cc in ws[1]:
                cc.font = Font(bold=True, color="FFFFFF")
                cc.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.3), 80)
            ws.freeze_panes = "A2"
    print(f"\n完成：{path}\n")
    print(f"{'訊號':<26}{'筆數':>8}{'落後根數':>8}{'已走ATR':>9}{'已走比例':>9}{'之後順向':>9}{'之後逆向':>9}{'順>逆%':>8}{'12根報酬':>10}")
    for _, r in S.iterrows():
        print(f"{r['訊號']:<26}{r['筆數']:>8}{r['中位數_落後根數']:>8.0f}{r['中位數_已走幅度ATR']:>9.2f}"
              f"{r['中位數_已走比例']:>9.0%}{r['中位數_之後最大順向ATR']:>9.2f}{r['中位數_之後最大逆向ATR']:>9.2f}"
              f"{r['順向>逆向比例']:>8.1f}{r['平均_之後12根報酬ATR']:>+10.3f}")
    del base


if __name__ == "__main__":
    main()
