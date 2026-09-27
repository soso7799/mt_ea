"""
H1 順勢（規則單純化）—— 小時線趨勢方向 + M5 進場

H1 趨勢（只用已收盤的 H1 K 線）：
  多頭 = 收盤 > EMA20 且 EMA20 > EMA50 且 EMA50 比 5 根前高；空頭反之；其他 = 無趨勢（不做）
事先定好的 3 種做法（參數不優化）：
  A 趨勢持有：H1 轉為多頭/空頭時下一根進場，趨勢結束才出場；安全停損 2 H1 ATR
  B 順勢回檔：H1 多頭時，M5 回檔碰到 M5 EMA20 後收回其上且為陽線 → 做多（空頭反之）；
     停損 = 最近 6 根 M5 低點外 0.1 H1 ATR（至少 0.3、至多 1.5 H1 ATR）；停利 2R；H1 趨勢結束也出場
  C 順勢回檔、讓利潤跑：進場與停損同 B，不設停利，H1 趨勢結束才出場
對照：B 的進場點改成「逆 H1 方向」做 —— 若 H1 方向真的有用，順勢應明顯好於逆勢
成交：M5 下一根開盤；買在 Ask、賣在 Bid；扣點差與外匯手續費；同商品一次一筆；同根碰到停損停利算停損
FTMO 試算：每筆風險 0.5% 帳戶，各商品合計的最大回撤與最差單日

用法：python h1_trend.py --m1-cache H:\\...\\export\\bars --out H:\\...\\reports [--years 2]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "ml"))

from build_report import DEFAULT_SYMBOLS      # noqa: E402
from indicators import atr, ema               # noqa: E402
from levels_study import is_fx                # noqa: E402
import optimize_wf as ow                      # noqa: E402

RISK_PCT = 0.5
VARIANTS = ["A 趨勢持有", "B 順勢回檔 停利2R", "C 順勢回檔 讓利潤跑", "對照：B 逆勢做"]


def h1_state(m1, m5_time):
    h1 = ow.resample(m1, 60)
    c = h1["close"].to_numpy(float)
    e20, e50 = ema(c, 20), ema(c, 50)
    a = atr(h1["high"].to_numpy(float), h1["low"].to_numpy(float), c, 14)
    e50_5 = np.r_[np.full(5, np.nan), e50[:-5]]
    st = np.where((c > e20) & (e20 > e50) & (e50 > e50_5), 1, np.where((c < e20) & (e20 < e50) & (e50 < e50_5), -1, 0))
    T = pd.DataFrame({"t": (h1["time"] + pd.Timedelta(minutes=60)).astype("datetime64[ns]"), "st": st, "a": a})
    left = pd.DataFrame({"t": pd.to_datetime(pd.Series(m5_time)).astype("datetime64[ns]").to_numpy()})
    m = pd.merge_asof(left, T, on="t", direction="backward")     # 只用 M5 開盤前已收盤的 H1
    return m["st"].fillna(0).to_numpy(int), m["a"].to_numpy(float)


def next_change(st):
    """每根之後第一個 st 與本根不同的索引（沒有則 n）。"""
    n = len(st)
    chg = np.flatnonzero(st[1:] != st[:-1]) + 1
    pos = np.searchsorted(chg, np.arange(n), side="right")
    return np.where(pos < len(chg), chg[np.minimum(pos, len(chg) - 1)], n)


def simulate(entries, o, h, l, c, sp, st, nxt, comm_px, use_flip=True):
    """entries: list of (e, d, sl, tp_or_nan)，於 e 開盤進場。回傳每筆 dict。"""
    n = len(c)
    out, busy = [], -1
    for e, d, sl, tp in entries:
        if e <= busy or e >= n - 1:
            continue
        entry = o[e] + (sp[e] if d > 0 else 0.0)
        dist = abs(entry - sl)
        if not dist > 0 or (sl - entry) * d >= 0:
            continue
        end = min(nxt[e] if use_flip and st[e] == d else n - 1, n - 1)
        if not use_flip:
            end = n - 1
        j = np.arange(e, end + 1)
        adj = 0.0 if d > 0 else sp[j]
        hi, lo = h[j] + adj, l[j] + adj
        s_hit = (lo <= sl) if d > 0 else (hi >= sl)
        t_hit = np.zeros(len(j), bool) if np.isnan(tp) else ((hi >= tp) if d > 0 else (lo <= tp))
        hit = s_hit | t_hit
        if hit.any():
            k = int(np.argmax(hit))
            op = o[j[k]] + (0.0 if d > 0 else sp[j[k]])
            if s_hit[k]:
                px = op if (op - sl) * d < 0 else sl        # 跳空穿過停損：以開盤價出場
                why = "停損"
            else:
                px = tp
                why = "停利"
            x = j[k]
        elif end < n - 1:
            x = end                                         # 趨勢結束：下一根開盤出場
            px = o[x] + (0.0 if d > 0 else sp[x])
            why = "趨勢結束"
        else:
            x = n - 1
            px = c[x] + (0.0 if d > 0 else sp[x])
            why = "資料結束"
        R = ((px - entry) * d - comm_px) / dist
        R0 = R + (comm_px + (sp[e] if d > 0 else sp[x])) / dist      # 約略加回點差與手續費
        out.append(dict(e=e, x=x, d=d, R=R, R0=R0, 出場=why, 持有根數=x - e))
        busy = x
    return out


def run_symbol(sym, m1, meta, commission):
    df = ow.resample(m1, 5)
    o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
    n = len(c)
    point = float(meta.get("point", 0.0001))
    contract = float(meta.get("contract") or 100000)
    sp = df["spread"].to_numpy(float) * point
    med = np.nanmedian(sp[sp > 0]) if (sp > 0).any() else point
    sp = np.where(sp > 0, sp, med)
    comm_px = commission / contract * (1.0 if sym.endswith("USD") else np.nanmedian(c)) if is_fx(sym) and commission else 0.0
    st, A = h1_state(m1, df["time"])
    nxt = next_change(st)
    e20 = ema(c, 20)

    # A：趨勢開始
    on = np.flatnonzero((st != 0) & (np.r_[0, st[:-1]] != st))
    ent_A = [(i, st[i], o[i] - st[i] * 2 * A[i], np.nan) for i in on if A[i] > 0]
    # B/C：回檔進場（訊號 K 收盤，下一根開盤進場）
    lo3 = pd.Series(l - e20).rolling(3).min().to_numpy()
    hi3 = pd.Series(h - e20).rolling(3).max().to_numpy()
    lo6 = pd.Series(l).rolling(6).min().to_numpy()
    hi6 = pd.Series(h).rolling(6).max().to_numpy()
    longs = (st == 1) & (lo3 <= 0) & (c > e20) & (c > o)
    shorts = (st == -1) & (hi3 >= 0) & (c < e20) & (c < o)
    ent_B, ent_C, ent_X = [], [], []
    for i in np.flatnonzero(longs | shorts):
        if i + 1 >= n or not A[i] > 0:
            continue
        d = 1 if longs[i] else -1
        e = i + 1
        if st[e] != d:
            continue
        ref = o[e]
        raw = (lo6[i] - 0.1 * A[i]) if d > 0 else (hi6[i] + 0.1 * A[i])
        dist = np.clip(abs(ref - raw), 0.3 * A[i], 1.5 * A[i])
        sl = ref - d * dist
        ent_B.append((e, d, sl, ref + d * 2 * dist))
        ent_C.append((e, d, sl, np.nan))
        ent_X.append((e, -d, ref + d * dist, ref - d * 2 * dist))
    res = {}
    res[VARIANTS[0]] = simulate(ent_A, o, h, l, c, sp, st, nxt, comm_px)
    res[VARIANTS[1]] = simulate(ent_B, o, h, l, c, sp, st, nxt, comm_px)
    res[VARIANTS[2]] = simulate(ent_C, o, h, l, c, sp, st, nxt, comm_px)
    res[VARIANTS[3]] = simulate(ent_X, o, h, l, c, sp, st, nxt, comm_px, use_flip=False)
    t = df["time"].to_numpy()
    rows = []
    for v, lst in res.items():
        for r in lst:
            rows.append(dict(商品=sym, 做法=v, 進場時間=t[r["e"]], 出場時間=t[r["x"]], 方向="多" if r["d"] > 0 else "空",
                             R=r["R"], 不含成本R=r["R0"], 出場=r["出場"], 持有小時=r["持有根數"] * 5 / 60))
    trend_share = (st != 0).mean()
    return rows, trend_share


def stat(R):
    R = np.asarray(R, float)
    n = len(R)
    if n < 2:
        return n, (R.mean() if n else np.nan), np.nan
    sd = R.std(ddof=1)
    return n, R.mean(), (R.mean() / sd * np.sqrt(n) if sd > 1e-9 else np.nan)


def ftmo(T):
    """每筆 0.5% 風險、以出場時間累計的最大回撤與最差單日（%）。"""
    if not len(T):
        return np.nan, np.nan, np.nan
    s = T.sort_values("出場時間")
    eq = 100 + np.cumsum(s["R"].to_numpy() * RISK_PCT)
    dd = (np.maximum.accumulate(np.r_[100, eq])[1:] - eq).max()
    day = s.groupby(pd.to_datetime(s["出場時間"]).dt.date)["R"].sum() * RISK_PCT
    return dd, day.min(), eq[-1] - 100


def main():
    ap = argparse.ArgumentParser(description="H1 順勢（規則單純化）")
    ap.add_argument("--symbols", nargs="*", default=DEFAULT_SYMBOLS)
    ap.add_argument("--m1-cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--years", type=float, default=2.0)
    ap.add_argument("--to", default=None)
    ap.add_argument("--commission", type=float, default=5.0)
    args = ap.parse_args()

    from mt5data import fetch_m1
    end = dt.datetime.fromisoformat(args.to) if args.to else dt.datetime.combine(dt.date.today(), dt.time())
    start = end - dt.timedelta(days=round(args.years * 365.25))
    rows, share = [], {}
    for s in args.symbols:
        try:
            got = fetch_m1([s], start - dt.timedelta(days=10), end, args.m1_cache, args.terminal)
        except SystemExit as e:
            print(f"⚠️ {s} 略過：{e}")
            continue
        m1, meta = got[s]
        m1 = m1[(m1["time"] >= pd.Timestamp(start)) & (m1["time"] < pd.Timestamp(end))]
        if len(m1) < 50000:
            continue
        r, sh = run_symbol(s, m1, meta, args.commission)
        rows += r
        share[s] = sh
        print(f"  {s}: 完成（H1 有趨勢的時間 {sh:.0%}）")
    if not rows:
        raise SystemExit("沒有資料")
    T = pd.DataFrame(rows)
    mid = T["進場時間"].quantile(0.5)
    summ, per = [], []
    for v in VARIANTS:
        g = T[T["做法"] == v]
        n, mu, tt = stat(g["R"])
        h1 = g.loc[g["進場時間"] < mid, "R"].mean()
        h2 = g.loc[g["進場時間"] >= mid, "R"].mean()
        bysym = g.groupby("商品")["R"].sum()
        win, loss = g.loc[g["R"] > 0, "R"].sum(), -g.loc[g["R"] < 0, "R"].sum()
        dd, worst, net = ftmo(g)
        ok = not v.startswith("對照") and n >= 100 and mu > 0 and (tt or 0) >= 2 and h1 > 0 and h2 > 0 and (bysym > 0).mean() >= 0.6
        summ.append(dict(做法=v, 筆數=n, 平均R=mu, t值=tt, 不含成本平均R=g["不含成本R"].mean(), 勝率=(g["R"] > 0).mean() * 100,
                         獲利因子=win / loss if loss > 0 else np.nan, 總R=g["R"].sum(), 前半平均R=h1, 後半平均R=h2,
                         獲利商品比例=(bysym > 0).mean() * 100, 平均持有小時=g["持有小時"].mean(),
                         **{"帳戶報酬%(每筆0.5%)": net, "最大回撤%": dd, "最差單日%": worst},
                         判定=("✅ 可用" if ok else ("對照" if v.startswith("對照") else "❌"))))
        for s, gs in g.groupby("商品"):
            n2, m2, t2 = stat(gs["R"])
            per.append(dict(商品=s, 做法=v, 筆數=n2, 平均R=m2, t值=t2, 總R=gs["R"].sum(), 勝率=(gs["R"] > 0).mean() * 100,
                            H1有趨勢時間比例=share.get(s, np.nan) * 100))
    S = pd.DataFrame(summ)
    Pm = pd.DataFrame(per)
    ym = T.assign(月=pd.to_datetime(T["出場時間"]).dt.to_period("M").astype(str)).pivot_table(
        index="月", columns="做法", values="R", aggfunc="sum").reset_index()
    ex = T.groupby(["做法", "出場"]).agg(筆數=("R", "size"), 平均R=("R", "mean")).reset_index()

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"H1順勢_{updated:%Y%m%d_%H%M}.xlsx")
    notes = [f"產生時間 {updated}；期間 {start.date()} ~ {end.date()}；商品 {T['商品'].nunique()} 個",
             "H1 趨勢：收盤 > EMA20 > EMA50 且 EMA50 上升 = 多頭；反之空頭；其餘不做。只用已收盤 H1。",
             "A 趨勢持有：趨勢開始進場、結束出場，安全停損 2 H1 ATR。",
             "B 順勢回檔：H1 多頭中 M5 回檔碰 EMA20 後收回其上的陽線 → 下一根開盤做多（空頭反之）；停損最近 6 根低點外、停利 2R；H1 趨勢結束也出場。",
             "C：同 B 但不設停利，H1 趨勢結束才出場。對照：B 的進場點反向做（逆 H1 方向）。",
             "成本：點差 + 外匯手續費 $5/手；『不含成本平均R』看規則本身。1R = 停損距離。",
             "判定 ✅：≥100 筆、平均R>0、t≥2、前後半都賺、≥60% 商品獲利。參數全部事先定好、沒有優化。",
             f"FTMO 試算：每筆風險 {RISK_PCT}% 帳戶，所有商品合計；FTMO 限制最大回撤 10%、單日 5%。"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": notes}).to_excel(xw, sheet_name="說明", index=False)
        S.round(3).to_excel(xw, sheet_name="總結", index=False)
        Pm.round(3).to_excel(xw, sheet_name="各商品", index=False)
        ym.round(2).to_excel(xw, sheet_name="每月R", index=False)
        ex.round(3).to_excel(xw, sheet_name="出場原因", index=False)
        T.tail(20000).round(3).to_excel(xw, sheet_name="交易明細", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.sheets.values():
            for cc in ws[1]:
                cc.font = Font(bold=True, color="FFFFFF")
                cc.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.3), 60)
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("✅"):
                        cell.fill = PatternFill("solid", fgColor="C6EFCE")
    print(f"\n完成：{path}\n")
    for _, r in S.iterrows():
        print(f"  {r['做法']:<14} {r['筆數']:>6} 筆  平均 {r['平均R']:+.3f}R（不含成本 {r['不含成本平均R']:+.3f}R）  t={r['t值']:.2f}  "
              f"前半 {r['前半平均R']:+.3f} 後半 {r['後半平均R']:+.3f}  獲利商品 {r['獲利商品比例']:.0f}%  "
              f"最大回撤 {r['最大回撤%']:.1f}%  → {r['判定']}")


if __name__ == "__main__":
    main()
