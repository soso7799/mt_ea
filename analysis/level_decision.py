"""
關卡「會衝過還是會反轉」—— 用你的指標在接近關卡時判斷

關卡（每天用前一日資料算好）：昨高/低/收、昨亞/歐/美盤 高/低/收（伺服器時間 0–9 / 9–16 / 16–24，同 YesterdayHiL）、
      上週高/低/收、近期壓力/支撐（H1 最近一個已確認擺動高/低，左右 3 根）
觸碰事件（M5）：價格從關卡一側靠近並碰到關卡（±0.1 H1 ATR），每條關卡每天只記第一次。
觸碰當下記錄（全部以「朝關卡的方向」為正）：
  J 線（你的 3D-2K 與標準 3K-2D，9/3/3）、J 斜率、MACD 柱（12/26）與柱斜率、趨勢線突破（Swing3，最近 3 根）、
  EMA 排列（8/21、20/50、50/200）、價格離 EMA20、TAI（MA28/5）動能、成交量比（同 ExcelMonitor_VolumePanel：本根量 ÷ 含本根 20 根平均）、近 3 根量、靠近速度、收盤穿越幅度、
  今日已走幅度、關卡共振數、關卡種類、時段
結果：從下一根開盤起算，先往穿越方向走 0.5 H1 ATR = 突破；先往回走 0.5 H1 ATR = 反轉（最多 4 小時）
判斷模型：梯度提升樹，逐月滾動 —— 只用「該月之前」的事件訓練，預測該月（沒看過的資料）
交易：P(突破) ≥ 0.60 → 順勢追突破；P(突破) ≤ 0.40 → 反向做反轉；其餘不做。
      事先定好 SL 0.5 / TP 0.75 H1 ATR（1.5R），最多 4 小時；下一根開盤成交、Bid/Ask、扣點差與外匯手續費；
      同商品一次只持一筆；同根同時碰到 SL/TP 算 SL（保守）。

用法：python level_decision.py --m1-cache H:\\...\\export\\bars --out H:\\...\\reports [--years 1]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "ml"))

from build_report import DEFAULT_SYMBOLS                                   # noqa: E402
from indicators import atr, ema, tai                                       # noqa: E402
from levels_study import LEVEL_NAMES, SESS, daily_levels, is_fx, session_hours   # noqa: E402
import optimize_wf as ow                                                   # noqa: E402

TOL, CONFL, BRK = 0.1, 0.3, 0.5
MAX_HOLD = 48                       # M5 根數 = 4 小時
SL_M, TP_M = 0.5, 0.75              # 事先定好（H1 ATR 倍數）
SLTP_INFO = [(0.5, 0.5), (0.5, 1.0), (1.0, 1.5)]
P_HI, P_LO = 0.60, 0.40
MIN_TRAIN_MONTHS = 3
FEATS = ["J(3D-2K)", "J(3K-2D)", "J斜率", "MACD柱", "MACD柱斜率", "趨勢線突破同向", "趨勢線突破反向",
         "EMA20-50排列", "EMA8-21排列", "EMA50-200排列", "離EMA20", "TAI動能", "量比", "近3根量比", "靠近速度", "收盤穿越", "今日幅度",
         "共振數", "關卡_昨日", "關卡_時段", "關卡_上週", "關卡_近期擺動", "時段_亞", "時段_歐", "時段_美"]


def h1_series(m1, m5_time):
    """H1 ATR 與 H1 擺動高/低（只用已收盤、已確認的資料），對應到每根 M5。"""
    h1 = ow.resample(m1, 60)
    H, Lw, C = (h1[k].to_numpy(float) for k in ("high", "low", "close"))
    a = atr(H, Lw, C, 14)
    close_t = h1["time"] + pd.Timedelta(minutes=60)
    hi_p, lo_p = ow.swing_pivots(H, Lw, 3)
    conf_h = hi_p + 3
    conf_l = lo_p + 3
    sh_ser = pd.Series(np.nan, index=range(len(h1)))
    sl_ser = pd.Series(np.nan, index=range(len(h1)))
    sh_ser.iloc[conf_h[conf_h < len(h1)]] = H[hi_p[conf_h < len(h1)]]
    sl_ser.iloc[conf_l[conf_l < len(h1)]] = Lw[lo_p[conf_l < len(h1)]]
    T = pd.DataFrame({"t": close_t, "atr": a, "sw_h": sh_ser.ffill().to_numpy(), "sw_l": sl_ser.ffill().to_numpy()})
    # MT5 資料的時間可能是秒精度、重取樣後是微秒精度，合併前統一成 ns
    T["t"] = pd.to_datetime(T["t"]).astype("datetime64[ns]")
    left = pd.DataFrame({"t": pd.to_datetime(pd.Series(m5_time)).astype("datetime64[ns]").to_numpy()})
    m = pd.merge_asof(left, T, on="t", direction="backward")
    return m["atr"].to_numpy(float), m["sw_h"].to_numpy(float), m["sw_l"].to_numpy(float)


def build_events(sym, m1, meta, commission):
    df = ow.resample(m1, 5)
    o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
    tv = df["tick_volume"].to_numpy(float)
    t = df["time"]
    n = len(c)
    point = float(meta.get("point", 0.0001))
    contract = float(meta.get("contract") or 100000)
    sp = df["spread"].to_numpy(float) * point
    med = np.nanmedian(sp[sp > 0]) if (sp > 0).any() else point
    sp = np.where(sp > 0, sp, med)
    comm_px = 0.0
    if is_fx(sym) and commission:
        comm_px = commission / contract * (1.0 if sym.endswith("USD") else np.nanmedian(c))

    A, swh, swl = h1_series(m1, t)
    uh = session_hours(t)
    lv, day = daily_levels(df, uh)
    lv["近期壓力"] = swh
    lv["近期支撐"] = swl
    names = LEVEL_NAMES + ["近期壓力", "近期支撐"]
    Lmat = np.vstack([lv[k] for k in names])

    # 指標（你的設定）
    K, D = ow.kdj(h, l, c, 9, 3, 3)
    J1, J2 = 3 * D - 2 * K, 3 * K - 2 * D
    mh = ema(c, 12) - ema(c, 26)
    e20, e50 = ema(c, 20), ema(c, 50)
    e8, e21, e200 = ema(c, 8), ema(c, 21), ema(c, 200)
    col = tai(o, h, l, c, ma_period=28, tai_period=5)[1]
    tsec = t.to_numpy().astype("datetime64[s]").astype(np.int64)
    tb, ts = ow.trendline_break(o, h, l, c, tsec, 3)
    tb3 = pd.Series(tb).rolling(3, min_periods=1).max().to_numpy().astype(bool)
    ts3 = pd.Series(ts).rolling(3, min_periods=1).max().to_numpy().astype(bool)
    vavg = pd.Series(tv).rolling(20, min_periods=20).mean().to_numpy()   # 同 ExcelMonitor_VolumePanel：含本根的 20 根平均
    v3 = pd.Series(tv).rolling(3).mean().to_numpy()
    dser = pd.Series(day)
    day_hi = pd.Series(h).groupby(dser).cummax().to_numpy()
    day_lo = pd.Series(l).groupby(dser).cummin().to_numpy()

    rows = []
    cprev = np.r_[np.nan, c[:-1]]
    for k, name in enumerate(names):
        L = Lmat[k]
        tol = TOL * A
        cond = (h >= L - tol) & (l <= L + tol) & (np.abs(cprev - L) > tol) & ~np.isnan(L) & (A > 0)
        idx = np.flatnonzero(cond)
        idx = idx[(idx >= 60) & (idx < n - 2)]
        if not len(idx):
            continue
        key = pd.DataFrame({"i": idx, "d": day[idx], "v": np.round(L[idx] / point)})
        idx = key.drop_duplicates(["d", "v"])["i"].to_numpy()
        for i in idx:
            a = 1 if cprev[i] < L[i] else -1
            others = np.delete(Lmat[:, i], k)
            dist = np.abs(others - L[i])
            grp = ("關卡_上週" if name.startswith("上週") else "關卡_近期擺動" if name.startswith("近期") else
                   "關卡_時段" if any(name.startswith(f"昨{s}") for s, _, _ in SESS) else "關卡_昨日")
            ss = next((s for s, x0, x1 in SESS if x0 <= uh[i] < x1), "其他")
            r = {"商品": sym, "time": t.iloc[i], "i": i, "關卡": name, "方向": "壓力" if a > 0 else "支撐", "a": a,
                 "價位": L[i], "A": A[i],
                 "J(3D-2K)": (J1[i] - 50) * a, "J(3K-2D)": (J2[i] - 50) * a, "J斜率": (J1[i] - J1[i - 1]) * a,
                 "MACD柱": mh[i] / A[i] * a, "MACD柱斜率": (mh[i] - mh[i - 1]) / A[i] * a,
                 "趨勢線突破同向": float(tb3[i] if a > 0 else ts3[i]), "趨勢線突破反向": float(ts3[i] if a > 0 else tb3[i]),
                 "EMA20-50排列": (e20[i] - e50[i]) / A[i] * a, "EMA8-21排列": (e8[i] - e21[i]) / A[i] * a,
                 "EMA50-200排列": (e50[i] - e200[i]) / A[i] * a, "離EMA20": (c[i] - e20[i]) / A[i] * a,
                 "TAI動能": (1.0 if (col[i] == 1 and a > 0) or (col[i] == 2 and a < 0) else
                           (-1.0 if col[i] in (1, 2) else 0.0)),
                 "量比": tv[i] / vavg[i] if vavg[i] > 0 else np.nan,
                 "近3根量比": v3[i] / vavg[i] if vavg[i] > 0 else np.nan,
                 "靠近速度": (c[i] - c[i - 6]) / A[i] * a, "收盤穿越": (c[i] - L[i]) / A[i] * a,
                 "今日幅度": (day_hi[i] - day_lo[i]) / A[i],
                 "共振數": float(np.sum((dist <= CONFL * A[i]) & (dist > point))),
                 "關卡_昨日": 0.0, "關卡_時段": 0.0, "關卡_上週": 0.0, "關卡_近期擺動": 0.0,
                 "時段_亞": float(ss == "亞"), "時段_歐": float(ss == "歐"), "時段_美": float(ss == "美")}
            r[grp] = 1.0
            rows.append(r)
    E = pd.DataFrame(rows)
    if not len(E):
        return E
    E = E.sort_values("time").reset_index(drop=True)
    i = E["i"].to_numpy()
    a = E["a"].to_numpy()
    # 結果標籤：先走 0.5 H1 ATR 的方向
    j0 = i + 1
    ref = o[j0]
    Jm = j0[:, None] + np.arange(MAX_HOLD)[None, :]
    valid = Jm < n
    Jc = np.minimum(Jm, n - 1)
    b = BRK * E["A"].to_numpy()
    up = (h[Jc] >= (ref + b)[:, None]) & valid
    dn = (l[Jc] <= (ref - b)[:, None]) & valid
    brk = np.where(a[:, None] > 0, up, dn)
    rev = np.where(a[:, None] > 0, dn, up)
    anyh = brk | rev
    first = np.argmax(anyh, 1)
    rr = np.arange(len(i))
    has = anyh.any(1)
    fb, fr = brk[rr, first], rev[rr, first]
    E["結果"] = np.where(~has, "未決", np.where(fb & fr, "不明", np.where(fb, "突破", "反轉")))
    # 交易 R：追突破（d=a）與做反轉（d=-a）
    Aarr = np.full(n, np.nan)
    Aarr[i] = E["A"].to_numpy()
    for sl_m, tp_m in [(SL_M, TP_M)] + SLTP_INFO:
        for mode, d in (("突破", a), ("反轉", -a)):
            ii, e, x, R = ow.exits(i.copy(), d.copy(), o, h, l, c, sp, Aarr, sl_m, tp_m, comm_px)
            assert len(ii) == len(i)       # 事件都在 i < n-2 且 A > 0，全部有效、順序不變
            E[f"R_{mode}_{sl_m}_{tp_m}"] = R
            if (sl_m, tp_m) == (SL_M, TP_M):
                E[f"e_{mode}"] = e
                E[f"x_{mode}"] = x
    return E


def walk_forward_model(E):
    from sklearn.ensemble import HistGradientBoostingClassifier
    E = E.copy()
    E["P突破"] = np.nan
    month = E["time"].dt.to_period("M")
    months = sorted(month.unique())
    lab = E["結果"].isin(["突破", "反轉"])
    y = (E["結果"] == "突破").astype(int)
    for m in months[MIN_TRAIN_MONTHS:]:
        cut = m.to_timestamp()
        tr = lab & (E["time"] < cut - pd.Timedelta(hours=4))
        te = month == m
        if tr.sum() < 500 or te.sum() == 0:
            continue
        clf = HistGradientBoostingClassifier(max_depth=3, learning_rate=0.05, max_iter=200,
                                             l2_regularization=1.0, min_samples_leaf=50, random_state=0)
        clf.fit(E.loc[tr, FEATS].to_numpy(float), y[tr])
        E.loc[te, "P突破"] = clf.predict_proba(E.loc[te, FEATS].to_numpy(float))[:, 1]
    return E


def auc(y, p):
    y, p = np.asarray(y), np.asarray(p)
    pos, neg = p[y == 1], p[y == 0]
    if not len(pos) or not len(neg):
        return np.nan
    r = pd.Series(np.r_[pos, neg]).rank().to_numpy()
    return (r[:len(pos)].sum() - len(pos) * (len(pos) + 1) / 2) / (len(pos) * len(neg))


def pick_trades(E, which):
    """which: '全部追突破' / '全部做反轉' / '模型'；同商品一次一筆。回傳 (R 陣列, 時間, 商品)。"""
    out = []
    for s, g in E.groupby("商品"):
        g = g.sort_values("time")
        if which == "全部追突破":
            sel = [(r, "突破") for r in g.to_dict("records")]
        elif which == "全部做反轉":
            sel = [(r, "反轉") for r in g.to_dict("records")]
        else:
            sel = [(r, "突破" if r["P突破"] >= P_HI else "反轉") for r in g.to_dict("records")
                   if not np.isnan(r["P突破"]) and (r["P突破"] >= P_HI or r["P突破"] <= P_LO)]
        busy = -1
        for r, mode in sel:
            e_, x_ = r[f"e_{mode}"], r[f"x_{mode}"]
            R = r[f"R_{mode}_{SL_M}_{TP_M}"]
            if e_ < 0 or np.isnan(R) or e_ <= busy:
                continue
            busy = x_
            out.append((r["time"], s, mode, R))
    return pd.DataFrame(out, columns=["time", "商品", "做法", "R"])


def stat(R):
    R = np.asarray(R, float)
    n = len(R)
    if n < 2:
        return n, (R.mean() if n else np.nan), np.nan
    sd = R.std(ddof=1)
    return n, R.mean(), (R.mean() / sd * np.sqrt(n) if sd > 1e-6 else np.nan)


def univariate(E, cut):
    lab = E[E["結果"].isin(["突破", "反轉"])]
    base_is = (lab.loc[lab["time"] < cut, "結果"] == "突破").mean()
    base_oos = (lab.loc[lab["time"] >= cut, "結果"] == "突破").mean()
    rows = []
    for f in FEATS:
        x = lab[f]
        if x.nunique() <= 3:
            bins = sorted(x.dropna().unique())
            lab_b = x.map(lambda v: f"={v:g}")
        else:
            q = np.nanquantile(x[lab["time"] < cut], [0.2, 0.4, 0.6, 0.8])
            q = np.unique(q)
            edges = np.r_[-np.inf, q, np.inf]
            lab_b = pd.cut(x, edges).astype(str)
        for bname, g in lab.groupby(lab_b):
            gi, go = g[g["time"] < cut], g[g["time"] >= cut]
            pi, po = (gi["結果"] == "突破").mean(), (go["結果"] == "突破").mean()
            zi = (pi - base_is) / np.sqrt(base_is * (1 - base_is) / max(len(gi), 1))
            zo = (po - base_oos) / np.sqrt(base_oos * (1 - base_oos) / max(len(go), 1))
            stable = len(gi) >= 100 and len(go) >= 50 and np.sign(zi) == np.sign(zo) and abs(zi) >= 3 and abs(zo) >= 2
            rows.append(dict(指標=f, 區間=bname, 前70筆數=len(gi), 前70突破率=pi * 100, 後30筆數=len(go),
                             後30突破率=po * 100, 前70偏離z=zi, 後30偏離z=zo,
                             判定=("✅ 穩定偏" + ("突破" if zi > 0 else "反轉")) if stable else ""))
    return pd.DataFrame(rows), base_is, base_oos


def main():
    ap = argparse.ArgumentParser(description="關卡突破/反轉判斷（你的指標）")
    ap.add_argument("--symbols", nargs="*", default=DEFAULT_SYMBOLS)
    ap.add_argument("--m1-cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--years", type=float, default=1.0)
    ap.add_argument("--to", default=None)
    ap.add_argument("--commission", type=float, default=5.0)
    args = ap.parse_args()

    from mt5data import fetch_m1
    end = dt.datetime.fromisoformat(args.to) if args.to else dt.datetime.combine(dt.date.today(), dt.time())
    start = end - dt.timedelta(days=round(args.years * 365.25))
    allE = []
    for s in args.symbols:
        try:
            got = fetch_m1([s], start - dt.timedelta(days=10), end, args.m1_cache, args.terminal)
        except SystemExit as e:
            print(f"⚠️ {s} 略過：{e}")
            continue
        m1, meta = got[s]
        m1 = m1[(m1["time"] >= pd.Timestamp(start)) & (m1["time"] < pd.Timestamp(end))]
        if len(m1) < 50000:
            continue
        E = build_events(s, m1, meta, args.commission)
        print(f"  {s}: 觸碰事件 {len(E)}")
        if len(E):
            allE.append(E)
    if not allE:
        raise SystemExit("沒有事件")
    E = pd.concat(allE, ignore_index=True).sort_values("time").reset_index(drop=True)
    print("逐月滾動訓練判斷模型 ...")
    E = walk_forward_model(E)
    cut = E["time"].quantile(0.7)

    lab = E["結果"].isin(["突破", "反轉"])
    O = E[lab & E["P突破"].notna()]
    A_ = auc((O["結果"] == "突破").astype(int), O["P突破"])
    cal = []
    if len(O):
        O = O.assign(bin=pd.cut(O["P突破"], [0, 0.3, 0.4, 0.45, 0.5, 0.55, 0.6, 0.7, 1.0]))
        for b, g in O.groupby("bin", observed=True):
            cal.append(dict(模型預測P突破=str(b), 筆數=len(g), 預測平均=g["P突破"].mean() * 100,
                            實際突破率=(g["結果"] == "突破").mean() * 100))

    oos_start = E.loc[E["P突破"].notna(), "time"].min()
    Eo = E[E["time"] >= oos_start] if pd.notna(oos_start) else E.iloc[0:0]
    strat = []
    trades_all = {}
    for which in ("全部追突破", "全部做反轉", "模型"):
        T = pick_trades(Eo, which)
        trades_all[which] = T
        n, mu, tt = stat(T["R"])
        half = T["time"].quantile(0.5) if len(T) else None
        h1 = stat(T.loc[T["time"] < half, "R"])[1] if len(T) else np.nan
        h2 = stat(T.loc[T["time"] >= half, "R"])[1] if len(T) else np.nan
        pos_sym = (T.groupby("商品")["R"].sum() > 0).mean() * 100 if len(T) else np.nan
        ok = which == "模型" and n >= 100 and mu > 0 and (tt or 0) >= 2 and h1 > 0 and h2 > 0
        strat.append(dict(做法=which, 筆數=n, 平均R=mu, t值=tt, 勝率=(T["R"] > 0).mean() * 100 if n else np.nan,
                          總R=T["R"].sum() if n else 0, 前半平均R=h1, 後半平均R=h2, 獲利商品比例=pos_sym,
                          判定=("✅ 可用" if ok else ("❌" if which == "模型" else "對照"))))
    S = pd.DataFrame(strat)
    U, b_is, b_oos = univariate(E, cut)
    lvl = E[lab].groupby(["關卡"]).agg(筆數=("結果", "size"), 突破率=("結果", lambda x: (x == "突破").mean() * 100)).reset_index()
    per_sym = []
    Tm = trades_all["模型"]
    for s, g in Tm.groupby("商品"):
        n, mu, tt = stat(g["R"])
        per_sym.append(dict(商品=s, 筆數=n, 平均R=mu, t值=tt, 追突破筆數=(g["做法"] == "突破").sum(),
                            做反轉筆數=(g["做法"] == "反轉").sum(), 總R=g["R"].sum()))
    info = []
    for sl_m, tp_m in SLTP_INFO:
        for mode in ("突破", "反轉"):
            col_ = f"R_{mode}_{sl_m}_{tp_m}"
            r = Eo[col_].dropna()
            n, mu, tt = stat(r)
            info.append(dict(SLTP=f"{sl_m}/{tp_m}", 做法=f"全部{mode}（不限一筆）", 筆數=n, 平均R=mu, t值=tt))

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"關卡突破反轉判斷_{updated:%Y%m%d_%H%M}.xlsx")
    notes = [f"產生時間 {updated}；期間 {start.date()} ~ {end.date()}；M5；商品 {E['商品'].nunique()} 個；觸碰事件 {len(E)}",
             "關卡：昨高/低/收、昨亞/歐/美盤高/低/收（伺服器時間 0–9/9–16/16–24）、上週高/低/收、近期壓力/支撐（H1 擺動點）。",
             "結果：下一根開盤起，先往穿越方向走 0.5 H1 ATR = 突破；先往回 0.5 H1 ATR = 反轉；4 小時內都沒有 = 未決。",
             f"整體突破率：前 70% 期間 {b_is:.1%}，後 30% 期間 {b_oos:.1%}（關卡本身偏反轉時會 < 50%）。",
             "判斷依據（觸碰當下、以朝關卡方向為正）：J 線(兩種公式)、J 斜率、MACD 柱與斜率、趨勢線突破、EMA20/50、TAI、量比、靠近速度、關卡種類、共振、時段。",
             "模型：梯度提升樹，逐月滾動，只用該月之前的事件訓練 → 預測該月；統計全部是沒看過的資料。",
             f"模型判斷力 AUC = {A_:.3f}（0.5 = 跟猜的一樣；0.55 以上才算有用；0.6 以上算不錯）。",
             f"交易（事先定好）：P≥{P_HI:.2f} 追突破、P≤{P_LO:.2f} 做反轉；SL {SL_M} / TP {TP_M} H1 ATR；最多 4 小時；同商品一次一筆；扣點差與手續費。",
             "判定 ✅：模型交易 ≥100 筆、平均R>0、t≥2、前後半都賺。『全部追突破 / 全部做反轉』是對照組。",
             "『指標狀態 vs 突破率』：每個指標分 5 段，看哪一段突破率明顯偏高/偏低；前 70%（z≥3）與後 30%（z≥2）都顯著且同方向才標 ✅（同時段多條關卡事件相關，門檻從嚴）。"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": notes}).to_excel(xw, sheet_name="說明", index=False)
        S.round(3).to_excel(xw, sheet_name="總結", index=False)
        pd.DataFrame(cal).round(1).to_excel(xw, sheet_name="模型準不準", index=False)
        U.round(2).to_excel(xw, sheet_name="指標狀態vs突破率", index=False)
        lvl.round(1).sort_values("突破率").to_excel(xw, sheet_name="各關卡突破率", index=False)
        pd.DataFrame(per_sym).round(3).to_excel(xw, sheet_name="各商品", index=False)
        pd.DataFrame(info).round(3).to_excel(xw, sheet_name="其他停損停利", index=False)
        cols = ["商品", "time", "關卡", "方向", "價位", "結果", "P突破"] + FEATS
        E[cols].tail(20000).round(4).to_excel(xw, sheet_name="事件明細", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.sheets.values():
            for cc in ws[1]:
                cc.font = Font(bold=True, color="FFFFFF")
                cc.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.3), 80)
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("✅"):
                        cell.fill = PatternFill("solid", fgColor="C6EFCE")
    print(f"\n完成：{path}")
    print(f"整體突破率 {b_is:.1%} / {b_oos:.1%}   模型判斷力 AUC = {A_:.3f}（0.5 = 猜）")
    for r in strat:
        print(f"  {r['做法']}: {r['筆數']} 筆  平均 {r['平均R']:+.3f}R  t={r['t值'] if r['t值'] == r['t值'] else 0:.2f}  → {r['判定']}")
    st = U[U["判定"].str.startswith("✅")]
    print(f"  穩定的指標狀態：{len(st)} 個")
    for _, r in st.head(12).iterrows():
        print(f"    {r['指標']} {r['區間']}: 突破率 {r['前70突破率']:.0f}% / {r['後30突破率']:.0f}%  {r['判定']}")


if __name__ == "__main__":
    main()
