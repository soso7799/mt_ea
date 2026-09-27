"""
指標應用診斷：每個指標的訊號出現後，價格是「延續」還是「回頭」？

對 16 商品 × 5 週期 × 11 指標（標準參數，不最佳化）：
  事件
    翻轉：指標狀態由空翻多（做多）/ 由多翻空（做空）
    極端：進入超賣（做多）/ 超買（做空）；布林/Keltner 為跌破下軌（做多）/ 突破上軌（做空）
          —— 極端事件的「做多」代表「反向用法」的方向
    共識：11 指標 ≥7 或 ≥9 同向（剛成立）
  量測
    訊號 K 線收盤後、下一根開盤進場，第 h 根收盤（h = 1,3,5,10,20）的報酬，
    以訊號當時 ATR(14) 為單位（R），可跨商品合併；另算點差成本（同單位）。
  分盤勢：全部 / ADX≥25 趨勢盤 / ADX≤20 盤整盤；並分前半段、後半段檢查穩定性。
  判讀
    照訊號做有效：平均 R > 成本 且 t ≥ 2 且前後半段同為正
    反著做有效：平均 R < -成本 且 t ≤ -2 且前後半段同為負（照訊號相反方向做才賺）
    無效：其他

用法：python diagnose.py --cache H:\\...\\export\\bars_mtf --out H:\\...\\reports [--offline]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import votes as V                                                  # noqa: E402
from analyze import closed_bars                                    # noqa: E402
from build_report import DEFAULT_SYMBOLS                           # noqa: E402
from data import TFS, fetch_all                                    # noqa: E402
from indicators import adx, atr, rsi, stoch, williams_r, cci, psy, bollinger, keltner, shift   # noqa: E402
from rules import onset                                            # noqa: E402

H = [1, 3, 5, 10, 20]
STD = {"MA": dict(s=10, l=50), "RSI": dict(n=14, lo=30, hi=70), "KD": dict(k=14, d=3, s=3), "PSY": dict(n=12),
       "WR": dict(n=14), "MTM": dict(n=10), "MACD": dict(f=12, s=26, g=9), "BOLL": dict(n=20, k=2.0),
       "CCI": dict(n=14), "BIAS": dict(n=20), "KELTNER": dict(n=20, m=2.0)}
REGIMES = ["全部", "趨勢盤", "盤整盤"]


def extremes(ind, o, h, l, c):
    """回傳 (做多事件布林陣列, 做空事件布林陣列)：進入超賣=做多、進入超買=做空。"""
    p = STD[ind]
    if ind == "RSI":
        v = rsi(c, p["n"]); lo, hi = v < p["lo"], v > p["hi"]
    elif ind == "KD":
        v, _ = stoch(h, l, c, p["k"], p["d"], p["s"]); lo, hi = v < 20, v > 80
    elif ind == "WR":
        v = williams_r(h, l, c, p["n"]); lo, hi = v < -80, v > -20
    elif ind == "CCI":
        v = cci(h, l, c, p["n"]); lo, hi = v < -100, v > 100
    elif ind == "PSY":
        v = psy(c, p["n"]); lo, hi = v < 25, v > 75
    elif ind == "BOLL":
        _, up, dn = bollinger(c, p["n"], p["k"]); lo, hi = c < dn, c > up
    elif ind == "KELTNER":
        _, up, dn = keltner(h, l, c, p["n"], p["m"]); lo, hi = c < dn, c > up
    else:
        return None
    return onset(np.nan_to_num(lo).astype(bool)), onset(np.nan_to_num(hi).astype(bool))


class Acc:
    """累加統計（n、和、平方和、正報酬數、成本和），前後半段分開。"""
    def __init__(self):
        self.d = {}

    def add(self, key, r, cost, half):
        for hh in (0, 1, 2):          # 0=全部 1=前半 2=後半
            if hh and half != hh:
                continue
            k = key + (hh,)
            s = self.d.setdefault(k, np.zeros(5))
            s += [1, r, r * r, r > 0, cost]

    def add_many(self, key, rs, costs, halves):
        for hh in (0, 1, 2):
            m = np.ones(len(rs), bool) if hh == 0 else halves == hh
            if not m.any():
                continue
            k = key + (hh,)
            s = self.d.setdefault(k, np.zeros(5))
            x = rs[m]
            s += [len(x), x.sum(), (x * x).sum(), (x > 0).sum(), costs[m].sum()]


def measure(acc, tag, sym, tf, idx, dirs, o, c, a, cost_price, regime):
    n = len(c)
    ok = (idx + 1 < n) & ~np.isnan(a[idx]) & (a[idx] > 0)
    idx, dirs = idx[ok], dirs[ok]
    if not len(idx):
        return
    half = np.where(idx < n // 2, 1, 2)
    entry = o[idx + 1]
    cst = cost_price / a[idx]
    for hz in H:
        j = idx + hz
        m = j < n
        if not m.any():
            continue
        r = (c[j[m]] - entry[m]) * dirs[m] / a[idx[m]]
        for rg in REGIMES:
            rm = np.ones(m.sum(), bool) if rg == "全部" else (regime[idx[m]] == rg)
            if rm.any():
                acc.add_many((tf, *tag, rg, hz), r[rm], cst[m][rm], half[m][rm])
                if hz == 5 and rg == "全部":
                    acc.add_many(("SYM", sym, tf, *tag), r[rm], cst[m][rm], half[m][rm])


def run_symbol(acc, sym, frames, meta, server_now):
    point = float(meta.get("point", 0.0001))
    for tf in TFS:
        df = frames.get((sym, tf))
        if df is None or len(df) < 300:
            continue
        df = closed_bars(df, tf, server_now)
        o, h, l, c = (df[k].to_numpy(float) for k in ("open", "high", "low", "close"))
        spr = df["spread"].replace(0, np.nan).median()
        cost_price = (spr if not np.isnan(spr) else 1) * point
        a = atr(h, l, c, 14)
        ax, _, _ = adx(h, l, c, 14)
        regime = np.where(ax >= 25, "趨勢盤", np.where(ax <= 20, "盤整盤", "其他"))
        states = []
        for ind in V.ORDER:
            st, _ = V.compute(ind, STD[ind], o, h, l, c)
            states.append(st)
            lo_, so_ = onset(st == 1), onset(st == -1)
            idx = np.concatenate([np.flatnonzero(lo_), np.flatnonzero(so_)])
            d = np.concatenate([np.ones(lo_.sum(), int), -np.ones(so_.sum(), int)])
            measure(acc, (ind, "翻轉"), sym, tf, idx, d, o, c, a, cost_price, regime)
            ex = extremes(ind, o, h, l, c)
            if ex:
                idx = np.concatenate([np.flatnonzero(ex[0]), np.flatnonzero(ex[1])])
                d = np.concatenate([np.ones(ex[0].sum(), int), -np.ones(ex[1].sum(), int)])
                measure(acc, (ind, "極端"), sym, tf, idx, d, o, c, a, cost_price, regime)
        st = np.vstack(states)
        lv, sv = (st == 1).sum(0), (st == -1).sum(0)
        for k in (7, 9):
            lo_, so_ = onset(lv >= k), onset(sv >= k)
            idx = np.concatenate([np.flatnonzero(lo_), np.flatnonzero(so_)])
            d = np.concatenate([np.ones(lo_.sum(), int), -np.ones(so_.sum(), int)])
            measure(acc, ("共識", f"≥{k}/11"), sym, tf, idx, d, o, c, a, cost_price, regime)


def stats(s):
    n, sm, sq, pos, cs = s
    if n < 2:
        return n, np.nan, np.nan, np.nan, np.nan
    mean = sm / n
    var = max(sq / n - mean ** 2, 0) * n / (n - 1)
    t = mean / np.sqrt(var / n) if var > 0 else np.nan
    return int(n), mean, t, pos / n * 100, cs / n


def verdict(mean, t, cost, m1, m2):
    if np.isnan(mean) or np.isnan(t):
        return "樣本不足"
    if mean > cost and t >= 2 and m1 > 0 and m2 > 0:
        return "照訊號做有效"
    if mean < -cost and t <= -2 and m1 < 0 and m2 < 0:
        return "反著做有效"
    if t >= 2 and mean > 0:
        return "照訊號方向但不夠付成本" if mean <= cost else "照訊號方向（前後不穩定）"
    if t <= -2 and mean < 0:
        return "反向但不夠付成本" if mean >= -cost else "反向（前後不穩定）"
    return "無效"


def main():
    ap = argparse.ArgumentParser(description="指標應用診斷：訊號後價格延續或回頭")
    ap.add_argument("--symbols", nargs="*", default=DEFAULT_SYMBOLS)
    ap.add_argument("--cache", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--terminal", default=None)
    ap.add_argument("--offline", action="store_true")
    args = ap.parse_args()

    frames, metas = fetch_all(args.symbols, args.cache, args.terminal, offline=args.offline, max_age_min=24 * 60)
    syms = [s for s in args.symbols if (s, "D1") in frames]
    ticks = [pd.Timestamp(m["tick_time"]) for m in metas.values() if m.get("tick_time")]
    server_now = max(ticks) if ticks else None
    acc = Acc()
    for s in syms:
        print(f"診斷 {s} ...")
        run_symbol(acc, s, frames, metas.get(s, {}), server_now)

    rows, sym_rows = [], []
    for k, v in acc.d.items():
        if k[0] == "SYM" or k[-1] != 0:
            continue
        tf, ind, ev, rg, hz, _ = k
        n, mean, t, win, cost = stats(v)
        m1 = stats(acc.d.get(k[:-1] + (1,), np.zeros(5)))[1]
        m2 = stats(acc.d.get(k[:-1] + (2,), np.zeros(5)))[1]
        rows.append(dict(週期=tf, 指標=ind, 事件=ev, 盤勢=rg, 持有根數=hz, 樣本數=n, 平均R=mean, t值=t, 勝率=win,
                         成本R=cost, 前半平均R=m1, 後半平均R=m2, 判讀=verdict(mean, t, cost, m1, m2)))
    for k, v in acc.d.items():
        if k[0] != "SYM" or k[-1] != 0:
            continue
        _, sym, tf, ind, ev, _ = k
        n, mean, t, win, cost = stats(v)
        sym_rows.append(dict(商品=sym, 週期=tf, 指標=ind, 事件=ev, 樣本數=n, 平均R_5根=mean, t值=t, 勝率=win, 成本R=cost))
    D = pd.DataFrame(rows)
    S = pd.DataFrame(sym_rows)
    ind_order = {x: i for i, x in enumerate(V.ORDER + ["共識"])}
    D = D.sort_values(by=["指標", "事件", "盤勢", "週期", "持有根數"],
                      key=lambda s: s.map(ind_order) if s.name == "指標" else
                      (s.map({t: i for i, t in enumerate(TFS)}) if s.name == "週期" else
                       (s.map({r: i for i, r in enumerate(REGIMES)}) if s.name == "盤勢" else s)))

    # 總覽：持有 5 根、每個（指標, 事件, 盤勢）× 週期 的判讀與平均 R
    ov = D[D["持有根數"] == 5]
    piv = []
    for (ind, ev, rg), g in ov.groupby(["指標", "事件", "盤勢"], sort=False):
        row = {"指標": ind, "事件": ev, "盤勢": rg}
        for tf in TFS:
            x = g[g["週期"] == tf]
            if len(x):
                x = x.iloc[0]
                row[tf] = f"{x['判讀']}（{x['平均R']:+.3f}R, t={x['t值']:.1f}）"
        piv.append(row)
    P = pd.DataFrame(piv)

    updated = dt.datetime.now().replace(microsecond=0)
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, f"指標應用診斷_{updated:%Y%m%d_%H%M}.xlsx")
    write(path, P, D, S, updated, syms)
    print(f"\n完成：{path}")
    # 終端摘要
    for label in ("照訊號做有效", "反著做有效"):
        x = ov[(ov["判讀"] == label) & (ov["盤勢"] == "全部")]
        print(f"\n{label}（持有 5 根、全部盤勢）: {len(x)} 項")
        for _, r_ in x.sort_values("t值", key=abs, ascending=False).head(15).iterrows():
            print(f"  {r_['週期']:>4} {r_['指標']:<8}{r_['事件']:<6} 平均 {r_['平均R']:+.3f}R  t={r_['t值']:+.1f}  "
                  f"成本 {r_['成本R']:.3f}R  n={int(r_['樣本數'])}")


def write(path, P, D, S, updated, syms):
    from openpyxl.styles import Font, PatternFill
    colors = {"照訊號做有效": "C6EFCE", "反著做有效": "BDD7EE", "照訊號方向但不夠付成本": "FFEB9C", "反向但不夠付成本": "FFEB9C"}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"說明": [
            f"產生時間 {updated}；商品：{', '.join(syms)}",
            "問題：每個指標的訊號出現後，價格是延續（順勢用對）還是回頭（應該反向用）？",
            "翻轉：指標狀態由空翻多 = 做多、由多翻空 = 做空（順勢用法）。",
            "極端：進入超賣 / 跌破下軌 = 做多、進入超買 / 突破上軌 = 做空（反向用法）。",
            "共識：11 指標 ≥7 或 ≥9 同向剛成立時，順著共識方向。",
            "報酬：訊號 K 線收盤後下一根開盤進場，持有 h 根後的收盤；單位 R = 訊號當時 ATR(14)，可跨商品合併。",
            "成本R：點差換算成 ATR 單位。短週期的 ATR 小，成本佔比大。",
            "照訊號做有效 = 平均 R 大於成本、t ≥ 2、前後半段都為正 → 照上面定義的方向做有優勢。",
            "反著做有效 = 平均 R 小於負成本、t ≤ -2、前後半段都為負 → 要跟上面定義的方向反過來做才有優勢（= 用反了）。",
            "  例：『MACD 翻轉』反著做有效 = MACD 翻多後價格反而下跌，順勢追 MACD 是錯的用法。",
            "  例：『RSI 極端』照訊號做有效 = RSI 進入超賣後價格回升，超賣買 / 超買賣是對的用法。",
            "…但不夠付成本 = 統計上有方向，但幅度小於點差，實際交易仍會虧。",
            "參數用標準值（MA 10/50、RSI 14 30/70、KD 14/3/3、PSY 12、WR 14、MTM 10、MACD 12/26/9、布林 20/2、CCI 14、BIAS 20、Keltner 20/2），不做最佳化。",
            "趨勢盤 = ADX(14) ≥25；盤整盤 = ADX ≤20。"]}).to_excel(xw, sheet_name="說明", index=False)
        P.to_excel(xw, sheet_name="總覽（持有5根）", index=False)
        D.round(4).to_excel(xw, sheet_name="明細", index=False)
        S.round(4).to_excel(xw, sheet_name="各商品（持有5根）", index=False)
        for name in ("總覽（持有5根）", "明細", "各商品（持有5根）", "說明"):
            ws = xw.sheets[name]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(x.value or "")) for x in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, w * 1.6), 60)
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    v = str(cell.value or "")
                    for key, colr in colors.items():
                        if v.startswith(key):
                            cell.fill = PatternFill("solid", fgColor=colr)
                            break


if __name__ == "__main__":
    main()
